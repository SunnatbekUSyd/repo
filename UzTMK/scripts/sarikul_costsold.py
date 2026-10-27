"""Sarikul: show mining and processing costs as two lines in CostSold (applied to the user's edited file).

Rows Сырьё / Вспомогательные материалы become the mining and processing lines, the three zero energy rows
are removed (references to CostSold are shifted), and «Прочие производственные расходы» keeps the rest.
"""
import json
import re
import sys

import openpyxl
from openpyxl.styles import Font

src, out = sys.argv[1], sys.argv[2]
D = json.load(open('full_sarikul.json'))
mining, processing = D['varlines'][0], D['varlines'][1]
assert mining['name_ru'].startswith('Добыча') and processing['name_ru'].startswith('Переработка')
YEARS = range(2027, 2034)

wb = openpyxl.load_workbook(src)
ws = wb['CostSold']
labels = [ws.cell(r, 2).value.strip() for r in range(9, 14)]
assert labels == ['Сырьё', 'Вспомогательные материалы', 'Электроэнергия', 'Природный газ / топливо', 'Вода'], labels
assert ws.cell(18, 2).value.strip() == 'Прочие производственные расходы (без НДПИ)'
for r in range(9, 14):
    assert all(abs(ws.cell(r, c).value) < 0.5 for c in range(4, 11)), r

ws.cell(9, 2).value = '   Затраты на добычу руды (материалы, ГСМ, ВВ, электроэнергия)'
ws.cell(10, 2).value = '   Затраты на переработку руды (реагенты, электроэнергия и пр.)'
for i, y in enumerate(YEARS):
    c = 4 + i
    m, p = mining['cost'][str(y)] or 0.0, processing['cost'][str(y)] or 0.0
    ws.cell(9, c).value = m
    ws.cell(10, c).value = p
    rest = ws.cell(18, c).value - m - p
    assert rest > -0.5, (y, rest)
    ws.cell(18, c).value = rest

# remove the three zero energy rows (11-13) and shift every reference to CostSold rows below them
DEL_FIRST, N = 11, 3
heights = {r: ws.row_dimensions[r].height for r in range(DEL_FIRST + N, ws.max_row + 1)}
ws.delete_rows(DEL_FIRST, N)
for r, h in heights.items():
    ws.row_dimensions[r - N].height = h
for r in range(ws.max_row + 1, ws.max_row + N + 1):
    ws.row_dimensions[r].height = None

TOK = re.compile(r"(?<![A-Za-z0-9_.А-Яа-я])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(\$?[A-Z]{1,3}\$?)(\d+)(?![\d(])")


def shift(formula, own):
    def f(m):
        pre, col, row = m.group(1), m.group(2), int(m.group(3))
        target = pre is None and own or pre in ('CostSold!', "'CostSold'!")
        if not target or row < DEL_FIRST:
            return m.group(0)
        assert row >= DEL_FIRST + N, ('reference to a deleted row', formula)
        return f'{pre or ""}{col}{row - N}'
    # leave string literals untouched
    parts = re.split(r'("[^"]*")', formula)
    return ''.join(p if p.startswith('"') else TOK.sub(f, p) for p in parts)


changed = 0
for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('='):
                nv = shift(v, w.title == 'CostSold')
                if nv != v:
                    c.value = nv
                    changed += 1

# Labour: note under the energy table
lb = wb['Labour']
assert lb.cell(9, 2).value == 'ВСЕГО ЭНЕРГОРЕСУРСЫ' and lb.cell(10, 2).value is None
lb.cell(10, 2).value = ('Энергоресурсы (электроэнергия, ГСМ) входят в удельные затраты на добычу и переработку руды '
                        '(Приложения 7 и 11); отдельная разбивка в ТЭО не приведена.')
lb.cell(10, 2).font = Font(name='Arial', size=10, italic=True)
wb.save(out)
print('formulas shifted:', changed)
