"""Build the bank-format business plan (yearly) for one UzTMK project from the extracted FS-model data.

Layout follows the bank template «БП расчет Avto Climate Control» (Приложения 1–16 + ИНФОРМАЦИЯ + NPV),
converted to yearly columns: 1-й год = 2027 (инвестиционный / льготный период), 2–7 годы = 2028–2033.
All figures come from the FS model (sheet «Исх_данные» keeps the source row of every number) and from
the advance-payment file (loan split); everything else is a formula.
"""
import datetime as dt
import math
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

YEARS = list(range(2027, 2034))          # bank years 1..7
YC = {y: 4 + i for i, y in enumerate(YEARS)}  # bank sheets: D..J
SRC_YEARS = list(range(2026, 2034))       # Исх_данные: E..L
SC = {y: 5 + i for i, y in enumerate(SRC_YEARS)}
SRC = 'Исх_данные'

F_ARIAL = 'Arial'
THIN = Side(style='thin', color='A6A6A6')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_HEAD = PatternFill('solid', fgColor='D9E1F2')
FILL_TOT = PatternFill('solid', fgColor='F2F2F2')
FILL_KEY = PatternFill('solid', fgColor='FFF2CC')
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x";(0.00"x");"-"'

CLS_RU = {'Buildings & structures': 'Здания, сооружения и передаточные устройства',
          'Plant & machinery': 'Машины и оборудование',
          'Vehicles': 'Транспортные средства',
          'Furniture, IT & office': 'Мебель, оргтехника и офисное оборудование',
          'Mine development': 'Горно-капитальные работы и геологоразведка',
          'Other tangible': 'Прочие основные средства',
          'Land': 'Земля'}
EQUIP = {'Plant & machinery', 'Vehicles', 'Furniture, IT & office'}   # "P" of the EPC contract
EL_RU = {'Raw materials': 'Сырьё', 'Auxiliary materials': 'Вспомогательные материалы',
         'Electricity': 'Электроэнергия', 'Gas / fuel': 'Природный газ / топливо', 'Water': 'Вода',
         'Payroll': 'Заработная плата', 'Social tax': 'Социальный налог', 'Depreciation': 'Амортизация',
         'Repairs': 'Ремонт и техобслуживание', 'Other': 'Прочие производственные расходы'}
ENERGY = ('Electricity', 'Gas / fuel', 'Water')

# Advance-payment file «14.10.2026 … W …xlsx» (sheet «драйвер лойиҳалар»): EPC contract E/P/C, bn UZS
ADV = [  # key, name (uz), term, capacity, total cost mln USD, E, P, C (bn UZS), UZS/USD for contract (thous.), 1st-tranche base
    ('sarikul', 'Сарикўл вольфрам конининг Марказий қисмини ўзлаштириш (1-босқич)', '2027 йил декабрь',
     'Йилига 600 минг тн маъдан қайта ишлаш', 158, 116.0322968, 532.896, 1041.567375, 12.2, 'C'),
    ('tungsten', 'Вольфрам концентратини қайта ишлаш янги гидрометаллургия заводини қуриш (1-босқич)',
     '2027 йил ноябрь', 'Йилига 5000 тн вольфрам концентрати қайта ишлаш', 80, 25.272, 293.4, 353.99904, 12, 'C'),
    ('moly', 'Молибден куйинди саноат маҳсулотини қайта ишлаш бўйича янги гидрометаллургия заводини қуриш',
     '2027 йил декабрь', 'Йилига 5000 тн молибден куйиндиси (концентрати) қайта ишлаш', 65, 25.35648,
     251.13528, 442.311, 12, 'F'),
]
ADV_RATE = 11.838      # UZS/USD (thous.) used in the file for the 30% advance
LOAN_TOTAL = 78_000_000
BANK = 'Национальный банк внешнеэкономической деятельности Республики Узбекистан'
BANK_S = 'НБУ'
BORROWER = 'АО «Узбекский комбинат технологических металлов» (АО «Ўзбекистон ТМК», UzTMK)'
# Russian rendering of the advance-file rows (name, implementation term, capacity)
ADV_RU = {
    'sarikul': ('Освоение центральной части вольфрамового месторождения Сарикуль (1-я очередь)', 'декабрь 2027 г.',
                'Переработка 600 тыс. т руды в год'),
    'tungsten': ('Строительство нового гидрометаллургического завода по переработке вольфрамового концентрата (1-я очередь)',
                 'ноябрь 2027 г.', 'Переработка 5 000 т вольфрамового концентрата в год'),
    'moly': ('Строительство нового гидрометаллургического завода по переработке промышленного продукта — молибденового огарка',
             'декабрь 2027 г.', 'Переработка 5 000 т молибденового огарка (концентрата) в год'),
}


def clean(t):
    for a, b in (('—', '-'), ('–', '-'), ('−', '-'), ('×', 'x'), ('≥', '>='), ('•', '-'), ('…', '...')):
        t = t.replace(a, b)
    return t


def font(bold=False, color='000000', size=10, italic=False):
    return Font(name=F_ARIAL, bold=bold, color=color, size=size, italic=italic)


class Book:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    def sheet(self, name, widths=None):
        ws = self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        for col, w in (widths or {}).items():
            ws.column_dimensions[col].width = w
        return ws


def put(ws, ref, value, fmt=None, bold=False, fill=None, color='000000', border=True, wrap=False,
        align=None, italic=False, size=10):
    c = ws[ref] if isinstance(ref, str) else ws.cell(row=ref[0], column=ref[1])
    if type(c).__name__ == 'MergedCell':
        return c
    if isinstance(value, str):
        value = clean(value)
        if value in ('=0', '0'):
            value = 0
    c.value = value
    c.font = font(bold, color, size, italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical='center')
    return c


def title(ws, text, appendix=None, sub=None, last_col=11):
    put(ws, 'B2', text, bold=True, size=12, border=False)
    if appendix:
        put(ws, (2, last_col), appendix, bold=True, border=False, align='right')
    if sub:
        put(ws, 'B3', sub, italic=True, border=False, color='595959')


def year_header(ws, row, label='ПОКАЗАТЕЛИ', unit='Ед. изм.', total=None, first=4, extra_left=None):
    total = None
    cols = [(2, label), (3, unit)] + [(first + i, f'{i + 1} год' + (' (льготный период)' if i == 0 else ''))
                                      for i in range(len(YEARS))]
    for c, lab in cols:
        put(ws, (row, c), lab, bold=True, fill=FILL_HEAD, align='center', wrap=True)
        put(ws, (row + 1, c), None, fill=FILL_HEAD)
        ws.merge_cells(start_row=row, start_column=c, end_row=row + 1, end_column=c)
    if total:
        put(ws, (row, first + 7), total, bold=True, fill=FILL_HEAD, align='center')
        put(ws, (row + 1, first + 7), '1–7 годы', bold=True, fill=FILL_HEAD, align='center')
    ws.row_dimensions[row].height = 30


def yrow(ws, r, label, f, unit='USD', fmt=NUM, bold=False, fill=None, total='sum', indent=0, note=None):
    """Write one yearly row: f(y, col_letter) -> formula/value."""
    total, note = None, None
    put(ws, (r, 2), ('   ' * indent) + label, bold=bold, fill=fill, wrap=True)
    put(ws, (r, 3), unit, fill=fill, align='center')
    for y in YEARS:
        c = YC[y]
        put(ws, (r, c), f(y, L(c)), fmt=fmt, bold=bold, fill=fill)
    if total:
        tc = 4 + 7
        if total == 'sum':
            v = f'=SUM(D{r}:J{r})'
        elif total == 'avg':
            v = f'=AVERAGE(D{r}:J{r})'
        elif total == 'last':
            v = f'=J{r}'
        else:
            v = total
        put(ws, (r, tc), v, fmt=fmt, bold=bold, fill=fill)
    if note:
        ws.cell(row=r, column=2).comment = Comment(note, 'Бизнес-план')
    return r


class Src:
    """Sheet «Исх_данные»: every model input with its source row."""

    def __init__(self, ws, model_file):
        self.ws = ws
        self.r = 5
        self.rows = {}
        title(ws, 'ИСХОДНЫЕ ДАННЫЕ ИЗ ТЭО-МОДЕЛИ (в долл. США, если не указано иное)', last_col=12,
              sub=f'Источник: финансовая модель ТЭО «{model_file}». Значения перенесены без изменений '
                  '(синий шрифт); столбец «Источник» — лист и строка модели.')
        for i, h in enumerate(['Код', 'Показатель', 'Ед. изм.', 'Источник (лист!строка ТЭО-модели)']
                              + [str(y) for y in SRC_YEARS]):
            put(ws, (4, 1 + i), h, bold=True, fill=FILL_HEAD, align='center', wrap=True)
        ws.column_dimensions['A'].width = 16
        ws.column_dimensions['B'].width = 62
        ws.column_dimensions['C'].width = 11
        ws.column_dimensions['D'].width = 30
        for y in SRC_YEARS:
            ws.column_dimensions[L(SC[y])].width = 15
        ws.freeze_panes = 'E5'

    def section(self, text):
        put(self.ws, (self.r, 1), None, fill=FILL_TOT)
        put(self.ws, (self.r, 2), text, bold=True, fill=FILL_TOT)
        self.r += 1

    def add(self, key, label, unit, source, series=None, scalar=None, fmt=NUM, sign=1):
        ws, r = self.ws, self.r
        put(ws, (r, 1), key, color='595959')
        put(ws, (r, 2), label, wrap=True)
        put(ws, (r, 3), unit, align='center')
        put(ws, (r, 4), source, color='595959')
        if series is not None:
            for y in SRC_YEARS:
                v = series.get(str(y), 0.0) if isinstance(series, dict) else series(y)
                put(ws, (r, SC[y]), (v or 0.0) * sign, fmt=fmt, color='0000FF')
        else:
            put(ws, (r, SC[2026]), scalar, fmt=fmt, color='0000FF')
        self.rows[key] = r
        self.r += 1
        return r

    def ref(self, key, y=None, absolute=True):
        r = self.rows[key]
        col = L(SC[y]) if y is not None else L(SC[2026])
        return f"'{SRC}'!${col}${r}" if absolute else f"'{SRC}'!{col}{r}"


def build(D, key, out_path):
    s = D['series']
    meta = D['meta']
    bk = Book()
    adv = {a[0]: a for a in ADV}
    prods = D['products']
    main = prods[0]
    # reference ("full capacity") year: first year with 100% utilisation, otherwise the year of maximum utilisation
    util = {y: main['util'][str(y)] for y in YEARS}
    full = [y for y in YEARS if util[y] >= 0.999]
    REF = full[0] if full else max(YEARS, key=lambda y: (util[y], -y))
    fp = dt.date.fromisoformat(meta['first_prod'])
    months_y1 = 12 - fp.month + 1 if fp.year == 2027 else 12
    REFL = f'{YEARS.index(REF) + 1}-й год'

    # ------------------------------------------------------------------ sheets (order as in the bank file)
    ws_sum = bk.sheet('Резюме', {'A': 2, 'B': 58, 'C': 12, **{L(c): 15 for c in range(4, 12)}})
    ws_inf = bk.sheet('ИНФОРМАЦИЯ', {'A': 2, 'B': 44, 'C': 52, 'D': 16, 'E': 16, 'F': 16, 'G': 16, 'H': 16,
                                     'I': 16, 'J': 16, 'K': 16, 'L': 16, 'M': 16, 'N': 16, 'O': 16})
    ws_pc = bk.sheet('ProjCost', {'A': 2, 'B': 80, 'C': 16, 'D': 16, 'E': 16, 'F': 12, 'G': 4, 'H': 58})
    ws_fp = bk.sheet('FinPlan', {'A': 2, 'B': 58, 'C': 16, 'D': 16, 'E': 16, 'F': 13, 'G': 19})
    ws_sh = bk.sheet('Share of costs', {'A': 2, 'B': 58, **{L(c): 15 for c in range(3, 10)}})
    ws_dep = bk.sheet('Depreciate', {'A': 2, 'B': 50, 'C': 30, 'D': 17, 'E': 9, 'F': 9,
                                     **{L(c): 14 for c in range(7, 15)}})
    ws_pp = bk.sheet('ProdPlan', {'A': 2, 'B': 50, 'C': 11, **{L(c): 14 for c in range(4, 12)}})
    ws_sp = bk.sheet('SalePlan', {'A': 2, 'B': 50, 'C': 11, **{L(c): 15 for c in range(4, 12)}})
    ws_rm = bk.sheet('RawMat', {'A': 2, 'B': 58, 'C': 12, 'D': 15, 'E': 15, 'F': 16, 'G': 10, 'H': 42})
    ws_lb = bk.sheet('Labour', {'A': 2, 'B': 52, 'C': 13, 'D': 16, 'E': 16, 'F': 16, 'G': 16, 'H': 30})
    ws_tx = bk.sheet('Taxes', {'A': 2, 'B': 58, 'C': 11, **{L(c): 15 for c in range(4, 12)}})
    ws_ct = bk.sheet('CostTotal', {'A': 2, 'B': 58, 'C': 16, 'D': 12, 'E': 12, 'F': 16, 'G': 16, 'H': 12})
    ws_cs = bk.sheet('CostSold', {'A': 2, 'B': 58, 'C': 9, **{L(c): 15 for c in range(4, 12)}})
    ws_ln = bk.sheet('LOANS', {'A': 7, 'B': 44, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': 16, 'H': 16,
                               'I': 16, 'J': 8, 'K': 3, 'L': 3})
    ws_pl = bk.sheet('ProfLoss', {'A': 2, 'B': 58, 'C': 9, **{L(c): 15 for c in range(4, 12)}})
    ws_wc = bk.sheet('WorkCap', {'A': 2, 'B': 58, 'C': 14, **{L(c): 15 for c in range(4, 12)}})
    ws_cf = bk.sheet('CashFlow', {'A': 2, 'B': 62, 'C': 9, **{L(c): 15 for c in range(4, 12)}})
    ws_bp = bk.sheet('BreakPoint', {'A': 2, 'B': 72, 'C': 18, 'D': 40})
    ws_npv = bk.sheet('NPV_IRR', {'A': 2, 'B': 44, 'C': 10, **{L(c): 16 for c in range(4, 12)}})
    ws_chk = bk.sheet('Проверки', {'A': 2, 'B': 80, 'C': 16, 'D': 16, 'E': 12})
    src = Src(bk.sheet(SRC), meta['file'])

    # ------------------------------------------------------------------ Исх_данные
    src.section('1. ПАРАМЕТРЫ ТЭО-МОДЕЛИ')
    src.add('wacc', 'Ставка дисконтирования проекта (WACC), номинальная в USD', '%', D['src']['wacc'],
            scalar=meta['wacc'], fmt=PCT)
    src.add('capbc_life', 'Срок амортизации капитализированных процентов (IAS 23), лет', 'лет',
            D['src']['capbc_life'], scalar=meta['capbc_life'], fmt='0')
    src.add('months_y1', f'Месяцев эксплуатации в 2027 г. (начало производства {fp:%d.%m.%Y})', 'мес.',
            'F1 General §5.04', scalar=months_y1, fmt='0')
    for k, lab, u, f in [('dso', 'Дебиторская задолженность', 'дней', '0'), ('rm_days', 'Запасы сырья и материалов', 'дней', '0'),
                         ('wip_days', 'Незавершённое производство', 'дней', '0'), ('fg_days', 'Готовая продукция (мин. запас)', 'дней', '0'),
                         ('dpo', 'Кредиторская задолженность', 'дней', '0'), ('wage_lag', 'Задолженность по зарплате', 'дней', '0'),
                         ('vat_refund_lag', 'Срок возмещения НДС (операционный)', 'дней', '0'),
                         ('capex_vat_lag', 'Срок возмещения НДС по капвложениям', 'дней', '0'),
                         ('soc_rate', 'Ставка социального налога', '%', PCT)]:
        src.add(k, 'Оборотный капитал / ФОТ: ' + lab if k != 'soc_rate' else lab, u, D['src'].get(k, ''),
                scalar=meta.get(k) or 0, fmt=f)
    for k, lab, f in [('npv_model', 'NPV проекта по ТЭО-модели (12 лет, исходная структура финансирования)', NUM),
                      ('irr_model', 'IRR проекта по ТЭО-модели', PCT),
                      ('payback_model', 'Срок окупаемости по ТЭО-модели, лет от начала модели', '0.0'),
                      ('dpayback_model', 'Дисконтированный срок окупаемости по ТЭО-модели, лет', '0.0'),
                      ('pi_model', 'Индекс доходности по ТЭО-модели', '0.00')]:
        src.add(k, lab, '', D['src'][k], scalar=meta[k], fmt=f)

    src.section('2. КУРС, НАЛОГОВЫЕ СТАВКИ')
    src.add('fx', 'Средний курс UZS за 1 USD', 'сум', D['src']['fx_avg'], s['fx_avg'], fmt=NUM2)
    src.add('cit_rate', 'Ставка налога на прибыль', '%', D['src']['cit_rate'], s['cit_rate'], fmt=PCT)
    src.add('vat_rate', 'Ставка НДС', '%', D['src']['vat_rate'], s['vat_rate'], fmt=PCT)
    src.add('tax_adj', 'Налоговые корректировки к прибыли до налогообложения', 'USD', D['src']['tax_adj'], s['tax_adj'])

    src.section('3. ПРОИЗВОДСТВО И РЕАЛИЗАЦИЯ')
    for i, p in enumerate(prods):
        n = p['name_ru']
        src.add(f'cap{i}', f'{n}: проектная мощность в год', p['unit'], p['src_util'], scalar=p['capacity'], fmt=NUM1)
        src.add(f'util{i}', f'{n}: освоение мощности', '%', p['src_util'], p['util'], fmt=PCT)
        src.add(f'prod{i}', f'{n}: производство', p['unit'], p['src_prod'], p['prod'], fmt=NUM1)
        src.add(f'sexp{i}', f'{n}: продажи — экспорт', p['unit'], p['src_sales_exp'], p['sales_exp'], fmt=NUM1)
        src.add(f'sdom{i}', f'{n}: продажи — внутренний рынок', p['unit'], p['src_sales_dom'], p['sales_dom'], fmt=NUM1)
        src.add(f'close{i}', f'{n}: запас на конец года', p['unit'], p['src_close'], p['close'], fmt=NUM1)
        src.add(f'price{i}', f'{n}: цена реализации', f'USD/{p["unit"]}', p['src_price'], p['price'], fmt=NUM2)
        src.add(f'rev{i}', f'{n}: выручка (контроль)', 'USD', p['src_rev'], p['rev'])

    src.section('4. ОТЧЁТ О ПРИБЫЛЯХ И УБЫТКАХ ТЭО-МОДЕЛИ (затраты со знаком «+»)')
    for k, lab, sg in [('rev', 'Выручка', 1), ('cos', 'Себестоимость продаж', -1), ('cos_da', 'в т.ч. амортизация в себестоимости', -1),
                       ('sell', 'Коммерческие расходы', -1), ('ga', 'Общие и административные расходы', -1),
                       ('ooi', 'Прочие операционные доходы', 1), ('ooe', 'Прочие операционные расходы (вкл. налог на имущество, земельный, рентный, резерв)', -1),
                       ('ebit', 'Операционная прибыль (EBIT)', 1), ('fin', 'Финансовые расходы (кредит ТЭО-модели — в бизнес-плане заменён)', -1),
                       ('pbt', 'Прибыль до налогообложения', 1), ('tax', 'Расход по налогу на прибыль', -1),
                       ('np', 'Чистая прибыль', 1), ('da', 'Амортизация (итого)', 1), ('ebitda', 'EBITDA', 1),
                       ('capbc_dep', 'Амортизация капитализированных процентов кредита ТЭО-модели', 1),
                       ('usd_prod', 'Справочно: производственные затраты, номинированные в USD', 1),
                       ('rev_usd', 'Справочно: выручка от продукции с ценой в USD', 1)]:
        src.add(k, lab, 'USD', D['src'][k], s[k], sign=sg)

    src.section('5. СЕБЕСТОИМОСТЬ ПО ЭЛЕМЕНТАМ ЗАТРАТ (ТЭО, форма F4, до изменения остатков)')
    el_keys = {'Raw materials': 'el_raw', 'Auxiliary materials': 'el_auxiliary', 'Electricity': 'el_electricity',
               'Gas / fuel': 'el_gas', 'Water': 'el_water', 'Payroll': 'el_payroll', 'Social tax': 'el_social',
               'Depreciation': 'el_depreciation', 'Repairs': 'el_repairs', 'Other': 'el_other'}
    for el, k in el_keys.items():
        src.add(k, EL_RU[el], 'USD', D['src'][k], s[k])
    src.add('el_total', 'Итого производственная себестоимость (с амортизацией)', 'USD', D['src']['el_total'], s['el_total'])
    src.add('met', 'в т.ч. налог за пользование недрами (НДПИ), в составе прочих', 'USD', D['src']['met'], s['met'])
    src.add('var_cogs', 'Переменная себестоимость (ТЭО)', 'USD', D['src']['var_cogs'], s['var_cogs'])

    src.section('6. ОПЕРАЦИОННЫЕ (ПЕРИОДНЫЕ) РАСХОДЫ И ПРОЧИЕ СТАТЬИ (ТЭО, формы F5, F6, F10)')
    for k, lab, sg in [('pers_ga', 'Административный персонал (с соц. налогом)', 1), ('opex_fixed', 'Постоянные операционные расходы проекта', 1),
                       ('opex_var', 'Переменные операционные расходы (вкл. коммерческие)', 1), ('opex_ga_da', 'Амортизация административных активов', 1),
                       ('contingency', 'Резерв на непредвиденные операционные расходы', -1), ('prop_tax', 'Налог на имущество', 1),
                       ('land_tax', 'Земельный налог', 1), ('rent_tax', 'Специальный рентный налог', 1),
                       ('vat_out', 'НДС исходящий (с реализации на внутреннем рынке)', 1), ('vat_in_op', 'НДС входящий — операционный', 1),
                       ('vat_in_capex', 'НДС входящий — по капитальным вложениям', 1), ('vat_net', 'Чистый НДС (+ к уплате / − к возмещению)', 1)]:
        src.add(k, lab, 'USD', D['src'][k], s[k], sign=sg)

    src.section('7. ДЕНЕЖНЫЕ ПОТОКИ И ОБОРОТНЫЙ КАПИТАЛ (ТЭО, отчёт о движении денежных средств и баланс)')
    for k, lab in [('cgo', 'Денежные средства от операционной деятельности до налога на прибыль и процентов'),
                   ('d_inv', 'Изменение запасов'), ('d_rec', 'Изменение дебиторской задолженности'),
                   ('d_vat', 'Изменение НДС к возмещению'), ('d_pay', 'Изменение кредиторской задолженности'),
                   ('d_wag', 'Изменение задолженности по зарплате и налогам с ФОТ'),
                   ('bs_inv', 'Запасы (на конец года)'), ('bs_rec', 'Торговая дебиторская задолженность (на конец года)'),
                   ('bs_vat', 'НДС к возмещению (на конец года)'), ('bs_pay', 'Торговая кредиторская задолженность (на конец года)'),
                   ('bs_wag', 'Задолженность по зарплате и налогам с ФОТ (на конец года)')]:
        src.add(k, lab, 'USD', D['src'][k], s[k])
    src.add('capex', 'Капитальные вложения (с резервом)', 'USD', D['src']['capex'], s['capex'])

    src.section('8. КАПИТАЛЬНЫЕ ВЛОЖЕНИЯ ПО ОБЪЕКТАМ (ТЭО, форма F8): освоение и амортизация')
    for j, cx in enumerate(D['capex']):
        src.add(f'cx{j}', f'{cx["code"]} {cx["name_ru"]} — освоение ({CLS_RU.get(cx["cls"], cx["cls"])}; срок {cx["life"]:.0f} лет)',
                'USD', cx['src'], cx['phase'])
        src.add(f'cxd{j}', f'{cx["code"]} {cx["name_ru"]} — амортизация', 'USD', cx['src'], cx['dep'])

    src.section(f'9. ПЕРЕМЕННЫЕ ЗАТРАТЫ ПО СТАТЬЯМ (ТЭО, форма F4): потребность и стоимость (стоимость = затраты на ед. × выпуск)')
    for j, vl in enumerate(D['varlines']):
        src.add(f'vq{j}', f'{vl["name_ru"]} — годовая потребность', vl['unit'], vl.get('src_qty', vl['src']), vl.get('qty', {}), fmt=NUM1)
        src.add(f'vc{j}', f'{vl["name_ru"]} — стоимость ({vl["ccy"]}; {EL_RU[vl["element"]]})', 'USD', vl['src'], vl['cost'])

    src.section('10. ПЕРСОНАЛ (ТЭО, форма F9): численность и фонд оплаты труда (брутто)')
    for j, ps in enumerate(D['positions']):
        src.add(f'hc{j}', f'{ps["name_ru"]} — численность', 'чел.', ps['src'], ps['headcount'], fmt=NUM1)
        src.add(f'mc{j}', f'{ps["name_ru"]} — месячная зарплата на 1 чел. ({ps["ccy"]})', ps['ccy'], ps['src'], ps['monthly_ccy'])
        src.add(f'pc{j}', f'{ps["name_ru"]} — годовой фонд оплаты труда', 'USD', ps['src'], ps['cost'])

    def X(k, y):
        return src.ref(k, y)

    # helpers for frequently used model items (positive = cost)
    nvar = len(D['varlines'])

    # ------------------------------------------------------------------ ИНФОРМАЦИЯ
    ws = ws_inf
    title(ws, 'ДАННЫЕ ПО ПРОЕКТУ', last_col=14)
    info = [('ЗАЕМЩИК', BORROWER), ('ИНН', '— (заполняется заемщиком)'),
            ('ДЕЯТЕЛЬНОСТЬ', 'Добыча и переработка редких и тугоплавких металлов (вольфрам, молибден)'),
            ('НАИМЕНОВАНИЕ ПРОЕКТА', ADV_RU[key][0] + f'\n({meta["name_ru"]})'),
            ('МЕСТО РЕАЛИЗАЦИИ', meta['site']),
            ('ПРОИЗВОДСТВЕННАЯ МОЩНОСТЬ', ADV_RU[key][2]),
            ('ОСНОВНАЯ ПРОДУКЦИЯ', ', '.join(p['name_ru'] for p in prods)),
            ('СТАДИЯ ПРОРАБОТКИ', {'Feasibility': 'ТЭО (технико-экономическое обоснование)',
                                   'Pre-feasibility': 'Предварительное ТЭО (Pre-FS)'}.get(meta['stage'], meta['stage'])),
            ('НАЧАЛО СТРОИТЕЛЬСТВА', dt.date.fromisoformat(meta['constr_start']).strftime('%d.%m.%Y')),
            ('НАЧАЛО ПРОИЗВОДСТВА', fp.strftime('%d.%m.%Y')),
            ('СРОК РЕАЛИЗАЦИИ (по данным заемщика)', ADV_RU[key][1])]
    r = 4
    for k, v in info:
        put(ws, (r, 2), k, bold=True)
        put(ws, (r, 3), v, wrap=True)
        ws.row_dimensions[r].height = 15 * max(1, sum(math.ceil(len(t) / 46) for t in str(v).split('\n')))
        r += 1
    r += 1
    put(ws, (r, 2), 'УЧАСТНИКИ ПРОЕКТА', bold=True, fill=FILL_HEAD); put(ws, (r, 3), None, fill=FILL_HEAD); r += 1
    for k, v in [('Заемщик (инициатор проекта)', BORROWER), ('Доля заемщика в проекте', '100%'),
                 ('Кредитор №1', BANK), ('Генеральный подрядчик', 'EPC-подрядчик (по EPC-контракту)')]:
        put(ws, (r, 2), k); put(ws, (r, 3), v, wrap=True)
        ws.row_dimensions[r].height = 30 if len(v) > 50 else 15
        r += 1
    r += 1
    put(ws, (r, 2), 'ВАЛЮТА ПРОЕКТА', bold=True); put(ws, (r, 3), 'Доллар США (экв.)'); r += 1
    put(ws, (r, 2), '1 USD = (сум, средний курс 1-го года)'); put(ws, (r, 3), f'={X("fx", 2027)}', fmt=NUM2); r_fx = r; r += 2

    put(ws, (r, 2), 'УСЛОВИЯ КРЕДИТА', bold=True, fill=FILL_HEAD); put(ws, (r, 3), None, fill=FILL_HEAD); r += 1
    r_loan = r
    terms = [('Сумма кредита, USD', None, NUM), ('Процентная ставка, % годовых', 0.07, PCT),
             ('Срок кредита, месяцев', 84, '0'), ('Льготный период по основному долгу, месяцев', 12, '0'),
             ('Инвестиционный период, месяцев', 12, '0'), ('Дата выдачи кредита', dt.date(2027, 1, 1), 'DD.MM.YYYY'),
             ('База начисления процентов, дней в году (по форме банка)', 360, '0'),
             ('Погашение основного долга', 'ежемесячно равными долями: 72 платежа, со 2-го по 7-й год', None),
             ('Уплата процентов', 'ежемесячно; в льготный период (1-й год) - за счёт собственных средств заемщика (включены в стоимость проекта)', None),
             ('Целевое назначение', '30% авансовый платёж по EPC-контракту (данные заемщика)', None),
             ('Резерв (непредвиденные расходы) — % от суммы кредита (по форме банка)', 0.01, PCT)]
    for k, v, f in terms:
        put(ws, (r, 2), k, wrap=True)
        c = put(ws, (r, 3), v, fmt=f, wrap=f is None, color='0000FF' if isinstance(v, (int, float, dt.date)) else '000000')
        lines = max(math.ceil(len(k) / 42), math.ceil(len(v) / 48) if isinstance(v, str) else 1)
        ws.row_dimensions[r].height = 15 * max(1, lines)
        r += 1
    INF = {'loan': f'ИНФОРМАЦИЯ!$C${r_loan}', 'rate': f'ИНФОРМАЦИЯ!$C${r_loan + 1}', 'term': f'ИНФОРМАЦИЯ!$C${r_loan + 2}',
           'grace': f'ИНФОРМАЦИЯ!$C${r_loan + 3}', 'inv': f'ИНФОРМАЦИЯ!$C${r_loan + 4}', 'date': f'ИНФОРМАЦИЯ!$C${r_loan + 5}',
           'basis': f'ИНФОРМАЦИЯ!$C${r_loan + 6}', 'reserve': f'ИНФОРМАЦИЯ!$C${r_loan + 10}', 'fx': f'ИНФОРМАЦИЯ!$C${r_fx}'}
    r += 1
    # allocation table of the 78 mln USD facility (advance-payment file)
    put(ws, (r, 2), 'РАСПРЕДЕЛЕНИЕ КРЕДИТА 78 МЛН USD МЕЖДУ ПРОЕКТАМИ (по 30% авансу EPC-контрактов, файл заемщика от 14.10.2026)',
        bold=True, border=False); r += 1
    hdr = ['Проект', 'Мощность', 'Стоимость проекта, млн USD', 'EPC: E (Engineering), млрд сум', 'EPC: P (Procurement), млрд сум',
           'EPC: C (Construction), млрд сум', 'EPC итого, млрд сум', 'Курс EPC (тыс. сум/USD)', 'EPC итого, экв. млн USD',
           '30% аванс, млрд сум', 'Курс аванса (тыс. сум/USD)', '30% аванс, млн USD', 'Доля в кредите', 'Кредит, USD']
    for i, h in enumerate(hdr):
        put(ws, (r, 2 + i), h, bold=True, fill=FILL_HEAD, wrap=True, align='center')
    ws.row_dimensions[r].height = 45
    r += 1
    r0 = r
    for a in ADV:
        cur = a[0] == key
        fill = FILL_KEY if cur else None
        put(ws, (r, 2), ADV_RU[a[0]][0], wrap=True, fill=fill, bold=cur)
        put(ws, (r, 3), ADV_RU[a[0]][2], wrap=True, fill=fill)
        put(ws, (r, 4), a[4], fmt=NUM, color='0000FF', fill=fill)
        put(ws, (r, 5), a[5], fmt=NUM2, color='0000FF', fill=fill)
        put(ws, (r, 6), a[6], fmt=NUM2, color='0000FF', fill=fill)
        put(ws, (r, 7), a[7], fmt=NUM2, color='0000FF', fill=fill)
        put(ws, (r, 8), f'=E{r}+F{r}+G{r}', fmt=NUM2, fill=fill)
        put(ws, (r, 9), a[8], fmt='0.0', color='0000FF', fill=fill)
        put(ws, (r, 10), f'=H{r}/I{r}', fmt=NUM2, fill=fill)
        put(ws, (r, 11), f'=H{r}*30%', fmt=NUM2, fill=fill)
        put(ws, (r, 12), ADV_RATE, fmt='0.000', color='0000FF', fill=fill)
        put(ws, (r, 13), f'=K{r}/L{r}', fmt=NUM2, fill=fill)
        put(ws, (r, 14), f'=M{r}/$M${r0 + 3}', fmt=PCT, fill=fill)
        put(ws, (r, 15), f'=N{r}*$O${r0 + 4}', fmt=NUM, fill=fill)
        ws.row_dimensions[r].height = 42
        if cur:
            r_cur = r
        r += 1
    put(ws, (r, 2), 'ИТОГО', bold=True, fill=FILL_TOT)
    for c in range(3, 16):
        col = L(c)
        v = None if c in (3, 9, 12) else f'=SUM({col}{r0}:{col}{r - 1})'
        put(ws, (r, c), v, fmt=PCT if c == 14 else (NUM if c in (4, 15) else NUM2), bold=True, fill=FILL_TOT)
    r += 1
    put(ws, (r, 14), 'Лимит кредита, USD', bold=True); put(ws, (r, 15), LOAN_TOTAL, fmt=NUM, color='0000FF', bold=True)
    ws[INF['loan'].split('!')[1].replace('$', '')].value = f'=O{r_cur}'
    ws[INF['loan'].split('!')[1].replace('$', '')].font = font(True)
    r += 2
    put(ws, (r, 2), 'Распределение кредита - пропорционально 30% авансу по EPC-контрактам, с приведением к лимиту 78,0 млн USD.', italic=True, border=False)

    # ------------------------------------------------------------------ LOANS (Приложение 12)
    ws = ws_ln
    title(ws, 'СХЕМА ВЫПЛАТЫ КРЕДИТА (в долл. США)', 'Приложение 12', last_col=8)
    params = [('Срок кредита', f'={INF["term"]}', 'месяцев', '0'), ('Льготный период', f'={INF["grace"]}', 'месяцев', '0'),
              ('Инвестиционный период', f'={INF["inv"]}', 'месяцев', '0'), ('Процентная ставка', f'={INF["rate"]}', 'годовых', PCT),
              ('Сумма кредита', f'={INF["loan"]}', 'USD', NUM), ('Дата выдачи', f'={INF["date"]}', '', 'DD.MM.YYYY'),
              ('База начисления процентов', f'={INF["basis"]}', 'дней', '0'),
              ('Количество платежей по основному долгу', '=C5-C6', 'платежей', '0'),
              ('Ежемесячный платёж по основному долгу', '=C9/C12', 'USD', NUM2)]
    for i, (k, v, u, f) in enumerate(params):
        put(ws, (5 + i, 2), k); put(ws, (5 + i, 3), v, fmt=f); put(ws, (5 + i, 4), u, border=False)
    hdr = ['№', 'Дата платежа', 'Остаток на начало', 'Погашение основного долга', 'Проценты', 'Всего выплат',
           'Остаток на конец', '№ года']
    for i, h in enumerate(hdr):
        put(ws, (15, 1 + i), h, bold=True, fill=FILL_HEAD, wrap=True, align='center')
    ws.row_dimensions[15].height = 30
    first = 16
    for k in range(1, 85):
        rr = first + k - 1
        put(ws, (rr, 1), k, fmt='0', align='center')
        put(ws, (rr, 2), f'=EOMONTH($C$10,{k - 1})', fmt='DD.MM.YYYY', align='center')
        put(ws, (rr, 3), '=$C$9' if k == 1 else f'=G{rr - 1}', fmt=NUM2)
        put(ws, (rr, 4), f'=IF(A{rr}>$C$6,MIN(C{rr},$C$13),0)', fmt=NUM2)
        put(ws, (rr, 5), f'=C{rr}*$C$8*DAY(B{rr})/$C$11', fmt=NUM2)
        put(ws, (rr, 6), f'=D{rr}+E{rr}', fmt=NUM2)
        put(ws, (rr, 7), f'=C{rr}-D{rr}', fmt=NUM2)
        put(ws, (rr, 8), f'=YEAR(B{rr})-YEAR($C$10)+1', fmt='0', align='center')
    last = first + 83
    rt = last + 1
    put(ws, (rt, 2), 'ИТОГО', bold=True, fill=FILL_TOT)
    for c in (4, 5, 6):
        put(ws, (rt, c), f'=SUM({L(c)}{first}:{L(c)}{last})', fmt=NUM2, bold=True, fill=FILL_TOT)
    # annual summary (horizontal)
    ra = rt + 3
    put(ws, (ra - 1, 2), 'ГОДОВАЯ СВОДКА ПО КРЕДИТУ', bold=True, border=False)
    put(ws, (ra, 2), 'ПОКАЗАТЕЛИ', bold=True, fill=FILL_HEAD)
    for i, y in enumerate(YEARS):
        put(ws, (ra, 3 + i), f'{i + 1} год', bold=True, fill=FILL_HEAD, align='center')
    LN = {}
    for j, (k, lab, colsrc) in enumerate([('drawn', 'Выдача кредита', None), ('princ', 'Погашение основного долга', 'D'),
                                          ('int', 'Проценты', 'E'), ('ds', 'Всего обслуживание долга', None),
                                          ('close', 'Остаток задолженности на конец года', None)]):
        rr = ra + 1 + j
        put(ws, (rr, 2), lab, bold=k in ('ds', 'close'))
        for i, y in enumerate(YEARS):
            c = 3 + i
            if k == 'drawn':
                v = '=$C$9' if i == 0 else 0
            elif k == 'ds':
                v = f'={L(c)}{ra + 2}+{L(c)}{ra + 3}'
            elif k == 'close':
                v = f'=$C$9-SUM($C${ra + 2}:{L(c)}{ra + 2})'
            else:
                v = f'=SUMIF($H${first}:$H${last},{i + 1},${colsrc}${first}:${colsrc}${last})'
            put(ws, (rr, c), v, fmt=NUM)
        LN[k] = rr

    def LNR(k, y):  # annual loan figure in LOANS sheet for bank year y
        return f'LOANS!${L(3 + YEARS.index(y))}${LN[k]}'
    I_GRACE = LNR('int', 2027)
    ws.freeze_panes = 'C16'

    # ------------------------------------------------------------------ Depreciate (Приложение 4)
    ws = ws_dep
    title(ws, 'АМОРТИЗАЦИЯ (в долл. США)', 'Приложение 4', last_col=14,
          sub='Линейный метод. Проценты льготного периода капитализированы в стоимость основных средств (МСФО IAS 23).')
    hdr = ['ОБЪЕКТЫ АМОРТИЗАЦИИ', 'Группа основных средств', 'Первоначальная стоимость', 'Срок, лет', 'Норма, %']
    for i, h in enumerate(hdr):
        put(ws, (5, 2 + i), h, bold=True, fill=FILL_HEAD, wrap=True, align='center')
    for i, y in enumerate(YEARS):
        put(ws, (5, 7 + i), f'{i + 1} год', bold=True, fill=FILL_HEAD, align='center')
    for c in range(2, 7 + len(YEARS)):
        put(ws, (6, c), None, fill=FILL_HEAD)
        ws.merge_cells(start_row=5, start_column=c, end_row=6, end_column=c)
    ws.row_dimensions[5].height = 30
    r = 7
    dep_first = r
    for j, cx in enumerate(D['capex']):
        put(ws, (r, 2), cx['name_ru'], wrap=True)
        put(ws, (r, 3), CLS_RU.get(cx['cls'], cx['cls']), wrap=True)
        put(ws, (r, 4), '=' + '+'.join(X(f'cx{j}', y) for y in SRC_YEARS), fmt=NUM)
        put(ws, (r, 5), cx['life'], fmt='0', color='0000FF')
        put(ws, (r, 6), f'=1/E{r}', fmt=PCT)
        for i, y in enumerate(YEARS):
            put(ws, (r, 7 + i), f'={X(f"cxd{j}", y)}', fmt=NUM)
        None
        r += 1
    r_capint = r
    put(ws, (r, 2), 'Капитализированные проценты по кредиту за льготный период (1-й год)', wrap=True)
    put(ws, (r, 3), 'Затраты по займам (МСФО IAS 23)')
    put(ws, (r, 4), f'={I_GRACE}', fmt=NUM)
    put(ws, (r, 5), f'={src.ref("capbc_life")}', fmt='0')
    put(ws, (r, 6), f'=1/E{r}', fmt=PCT)
    for i, y in enumerate(YEARS):
        v = f'=D{r}/E{r}*{months_y1}/12' if i == 0 and months_y1 < 12 else f'=D{r}/E{r}'
        put(ws, (r, 7 + i), v, fmt=NUM)
    None
    r += 1
    r_dep_tot = r
    put(ws, (r, 2), 'ВСЕГО', bold=True, fill=FILL_TOT); put(ws, (r, 3), 'Х', fill=FILL_TOT, align='center')
    for c in [4] + list(range(7, 14)):
        put(ws, (r, c), f'=SUM({L(c)}{dep_first}:{L(c)}{r - 1})', fmt=NUM, bold=True, fill=FILL_TOT)
    put(ws, (r, 5), None, fill=FILL_TOT); put(ws, (r, 6), None, fill=FILL_TOT)
    r += 1
    put(ws, (r, 2), 'в т.ч. в себестоимости продукции'); r_dep_cos = r
    for i, y in enumerate(YEARS):
        put(ws, (r, 7 + i), f'={L(7 + i)}{r_dep_tot}-{L(7 + i)}{r + 1}', fmt=NUM)
    None; r += 1
    put(ws, (r, 2), 'в т.ч. в административных расходах'); r_dep_ga = r
    for i, y in enumerate(YEARS):
        put(ws, (r, 7 + i), f'={X("opex_ga_da", y)}', fmt=NUM)
    None; r += 2
    put(ws, (r, 2), 'Капитальные вложения (освоение; 1-й год вкл. подготовительный период)', wrap=True); r_cx = r
    for i, y in enumerate(YEARS):
        yy = [2026, 2027] if y == 2027 else [y]
        put(ws, (r, 7 + i), '=' + '+'.join(X(f'cx{j}', z) for j in range(len(D['capex'])) for z in yy), fmt=NUM)
    r += 1
    put(ws, (r, 2), 'Накопленная амортизация', bold=True); r_acc = r
    for i in range(7):
        put(ws, (r, 7 + i), f'=SUM($G${r_dep_tot}:{L(7 + i)}{r_dep_tot})', fmt=NUM)
    r += 1
    put(ws, (r, 2), 'Остаточная стоимость основных средств (на конец года)', bold=True); r_nbv = r
    for i, y in enumerate(YEARS):
        c = L(7 + i)
        v = f'={c}{r_cx}+$D${r_capint}-{c}{r_dep_tot}' if i == 0 else f'={L(6 + i)}{r}+{c}{r_cx}-{c}{r_dep_tot}'
        put(ws, (r, 7 + i), v, fmt=NUM)
    DEP = {'tot': r_dep_tot, 'cos': r_dep_cos, 'ga': r_dep_ga, 'capint': r_capint, 'capex': r_cx}
    for rr in (r_dep_cos, r_dep_ga, r_cx, r_acc, r_nbv):
        for c in range(3, 7):
            put(ws, (rr, c), None)
        ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical='center')

    def DEPR(k, y):
        return f'Depreciate!${L(7 + YEARS.index(y))}${DEP[k]}'

    # ------------------------------------------------------------------ ProjCost (Приложение 1)
    ws = ws_pc
    title(ws, 'ПЕРВОНАЧАЛЬНАЯ СТОИМОСТЬ ПРОЕКТА (в экв. долл. США)', 'Приложение 1', last_col=6,
          sub='Инвестиционный период: с начала строительства до конца 1-го года (льготный период по кредиту).')
    for i, h in enumerate(['Наименование затрат', 'МЕСТНАЯ ВАЛЮТА', 'ТВЕРДАЯ ВАЛЮТА', 'ВСЕГО', 'СТРУКТУРА (%)']):
        put(ws, (4, 2 + i), h, bold=True, fill=FILL_HEAD, align='center', wrap=True)
    ws.row_dimensions[4].height = 30
    classes = []
    for cx in D['capex']:
        if cx['cls'] not in classes:
            classes.append(cx['cls'])
    r = 5
    PC = {}
    put(ws, (r, 2), 'ВСЕГО ФИКСИРОВАННЫЕ АКТИВЫ', bold=True, fill=FILL_TOT); r_fa = r; r += 1
    for cl in classes:
        idx = [j for j, cx in enumerate(D['capex']) if cx['cls'] == cl]
        usd = [j for j in idx if D['capex'][j]['ccy'] == 'USD']
        uzs = [j for j in idx if D['capex'][j]['ccy'] != 'USD']
        put(ws, (r, 2), CLS_RU.get(cl, cl))
        f_l = '=' + ('+'.join(f'{X(f"cx{j}", 2026)}+{X(f"cx{j}", 2027)}' for j in uzs) if uzs else '0')
        f_h = '=' + ('+'.join(f'{X(f"cx{j}", 2026)}+{X(f"cx{j}", 2027)}' for j in usd) if usd else '0')
        put(ws, (r, 3), f_l, fmt=NUM); put(ws, (r, 4), f_h, fmt=NUM)
        put(ws, (r, 5), f'=C{r}+D{r}', fmt=NUM)
        PC['cls_' + cl] = r
        r += 1
    r_fa_end = r - 1
    for c in (3, 4, 5):
        put(ws, (r_fa, c), f'=SUM({L(c)}{r_fa + 1}:{L(c)}{r_fa_end})', fmt=NUM, bold=True, fill=FILL_TOT)
    put(ws, (r, 2), 'РАБОЧИЙ КАПИТАЛ', bold=True, fill=FILL_TOT); r_wc = r
    put(ws, (r, 5), f'=C{r}+D{r}', fmt=NUM, bold=True, fill=FILL_TOT); r += 1
    put(ws, (r, 2), 'ПРОЧИЕ ФИНАНСОВЫЕ ИЗДЕРЖКИ, ВКЛЮЧАЯ:', bold=True, fill=FILL_TOT); r_fin = r; r += 1
    put(ws, (r, 2), '   Проценты по кредиту за льготный период (1-й год) - за счёт собственных средств')
    put(ws, (r, 3), 0, fmt=NUM); put(ws, (r, 4), f'={I_GRACE}', fmt=NUM); put(ws, (r, 5), f'=C{r}+D{r}', fmt=NUM)
    r_int = r; r += 1
    put(ws, (r, 2), '   Расходы подготовительного периода (административные и пусковые расходы)')
    put(ws, (r, 3), f'=-{X("ebitda", 2026)}', fmt=NUM); put(ws, (r, 4), 0, fmt=NUM); put(ws, (r, 5), f'=C{r}+D{r}', fmt=NUM)
    r_pre = r; r += 1
    for c in (3, 4, 5):
        put(ws, (r_fin, c), f'=SUM({L(c)}{r_fin + 1}:{L(c)}{r - 1})', fmt=NUM, bold=True, fill=FILL_TOT)
    put(ws, (r, 2), 'НЕПРЕДВИДЕННЫЕ РАСХОДЫ/РЕЗЕРВ (1% от суммы кредита)', bold=True, fill=FILL_TOT); r_res = r
    put(ws, (r, 3), 0, fmt=NUM, fill=FILL_TOT); put(ws, (r, 4), f'={INF["loan"]}*{INF["reserve"]}', fmt=NUM, fill=FILL_TOT, bold=True)
    put(ws, (r, 5), f'=C{r}+D{r}', fmt=NUM, bold=True, fill=FILL_TOT)
    r += 2
    r_pc_tot = r
    put(ws, (r, 2), 'ВСЕГО ПЕРВОНАЧАЛЬНАЯ СТОИМОСТЬ ПРОЕКТА', bold=True, fill=FILL_KEY)
    for c in (3, 4, 5):
        put(ws, (r, c), f'={L(c)}{r_fa}+{L(c)}{r_wc}+{L(c)}{r_fin}+{L(c)}{r_res}', fmt=NUM, bold=True, fill=FILL_KEY)
    for rr in range(r_fa, r + 1):
        if ws.cell(row=rr, column=5).value is not None:
            put(ws, (rr, 6), f'=IFERROR(E{rr}/$E${r_pc_tot},0)', fmt=PCT, bold=rr in (r_fa, r_wc, r_fin, r_res, r_pc_tot))
    r += 2
    PCR = {'fa': r_fa, 'wc': r_wc, 'fin': r_fin, 'int': r_int, 'pre': r_pre, 'res': r_res, 'tot': r_pc_tot}

    # ------------------------------------------------------------------ Share of costs (Приложение 3)
    ws = ws_sh
    title(ws, 'ПЕРВОНАЧАЛЬНАЯ СТОИМОСТЬ ПРОЕКТА ПО ИСТОЧНИКАМ ФИНАНСИРОВАНИЯ (в экв. долл. США)', 'Приложение 3', last_col=8)
    put(ws, (4, 2), 'Наименование затрат', bold=True, fill=FILL_HEAD, align='center'); put(ws, (5, 2), None, fill=FILL_HEAD)
    ws.merge_cells(start_row=4, start_column=2, end_row=5, end_column=2)
    ws.row_dimensions[4].height = 30
    for c0, h in [(3, 'Собственные средства заемщика'), (5, f'Кредит {BANK_S}'), (7, 'Всего')]:
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 1)
        put(ws, (4, c0), h, bold=True, fill=FILL_HEAD, align='center', wrap=True); put(ws, (4, c0 + 1), None, fill=FILL_HEAD)
        put(ws, (5, c0), 'в местной валюте', bold=True, fill=FILL_HEAD, align='center', wrap=True)
        put(ws, (5, c0 + 1), 'в твердой валюте', bold=True, fill=FILL_HEAD, align='center', wrap=True)
    ws.row_dimensions[5].height = 28
    # loan allocation: P share of EPC -> equipment classes, E+C share -> other classes (pro rata within group)
    advr = f'ИНФОРМАЦИЯ!$F${r_cur}/ИНФОРМАЦИЯ!$H${r_cur}'   # P / (E+P+C)
    eq_cls = [cl for cl in classes if cl in EQUIP]
    oth_cls = [cl for cl in classes if cl not in EQUIP]
    def pc_tot(cl):
        return f"ProjCost!$E${PC['cls_' + cl]}"
    r = 6
    SH = {}
    put(ws, (r, 2), 'ВСЕГО ФИКСИРОВАННЫЕ АКТИВЫ', bold=True, fill=FILL_TOT); r_sfa = r; r += 1
    for cl in classes:
        grp = eq_cls if cl in EQUIP else oth_cls
        share = advr if cl in EQUIP else f'(1-{advr})'
        if not grp:
            share = '1'
        denom = '+'.join(pc_tot(g) for g in grp)
        put(ws, (r, 2), CLS_RU.get(cl, cl))
        put(ws, (r, 5), 0, fmt=NUM)
        put(ws, (r, 6), f'={INF["loan"]}*{share}*{pc_tot(cl)}/({denom})', fmt=NUM)
        put(ws, (r, 3), f"=ProjCost!$C${PC['cls_' + cl]}", fmt=NUM)
        put(ws, (r, 4), f"=ProjCost!$D${PC['cls_' + cl]}-F{r}", fmt=NUM)
        put(ws, (r, 7), f'=C{r}+E{r}', fmt=NUM); put(ws, (r, 8), f'=D{r}+F{r}', fmt=NUM)
        r += 1
    for c in range(3, 9):
        put(ws, (r_sfa, c), f'=SUM({L(c)}{r_sfa + 1}:{L(c)}{r - 1})', fmt=NUM, bold=True, fill=FILL_TOT)
    for lab, pr in [('РАБОЧИЙ КАПИТАЛ', 'wc'), ('ПРОЧИЕ ФИНАНСОВЫЕ ИЗДЕРЖКИ', 'fin'),
                    ('   в т.ч. проценты по кредиту за льготный период', 'int'),
                    ('   в т.ч. расходы подготовительного периода', 'pre'), ('НЕПРЕДВИДЕННЫЕ РАСХОДЫ/РЕЗЕРВ', 'res')]:
        b = not lab.startswith('   ')
        put(ws, (r, 2), lab, bold=b, fill=FILL_TOT if b else None)
        put(ws, (r, 3), f'=ProjCost!$C${PCR[pr]}', fmt=NUM, bold=b, fill=FILL_TOT if b else None)
        put(ws, (r, 4), f'=ProjCost!$D${PCR[pr]}', fmt=NUM, bold=b, fill=FILL_TOT if b else None)
        put(ws, (r, 5), 0, fmt=NUM, fill=FILL_TOT if b else None); put(ws, (r, 6), 0, fmt=NUM, fill=FILL_TOT if b else None)
        put(ws, (r, 7), f'=C{r}+E{r}', fmt=NUM, bold=b, fill=FILL_TOT if b else None)
        put(ws, (r, 8), f'=D{r}+F{r}', fmt=NUM, bold=b, fill=FILL_TOT if b else None)
        SH[pr] = r
        r += 1
    r += 1
    r_sh_tot = r
    put(ws, (r, 2), 'ВСЕГО ПЕРВОНАЧАЛЬНАЯ СТОИМОСТЬ ПРОЕКТА', bold=True, fill=FILL_KEY)
    for c in range(3, 9):
        put(ws, (r, c), f'={L(c)}{r_sfa}+{L(c)}{SH["wc"]}+{L(c)}{SH["fin"]}+{L(c)}{SH["res"]}', fmt=NUM, bold=True, fill=FILL_KEY)
    r += 1
    put(ws, (r, 2), 'Итого по источнику', bold=True)
    put(ws, (r, 3), f'=C{r_sh_tot}+D{r_sh_tot}', fmt=NUM, bold=True); put(ws, (r, 5), f'=E{r_sh_tot}+F{r_sh_tot}', fmt=NUM, bold=True)
    put(ws, (r, 7), f'=G{r_sh_tot}+H{r_sh_tot}', fmt=NUM, bold=True)
    r_sh_src = r
    r += 2
    put(ws, (r, 2), f'Кредит {BANK_S} направляется на 30% авансовый платёж по EPC-контракту; распределение по группам активов — '
        'пропорционально составу EPC-контракта: P (поставка оборудования) — машины и оборудование, E+C (проектирование и '
        'строительство) — здания, сооружения и прочие объекты.', italic=True, border=False, wrap=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8); ws.row_dimensions[r].height = 42

    # ------------------------------------------------------------------ FinPlan (Приложение 2)
    ws = ws_fp
    title(ws, 'ФИНАНСОВЫЙ ПЛАН (в экв. долл. США)', 'Приложение 2', last_col=7)
    for i, h in enumerate(['ПОКАЗАТЕЛИ', 'МЕСТНАЯ ВАЛЮТА', 'ТВЁРДАЯ ВАЛЮТА', 'ВСЕГО', 'СТРУКТУРА (%)', 'ДОЛЯ В ФИНАНСИРОВАНИИ']):
        put(ws, (4, 2 + i), h, bold=True, fill=FILL_HEAD, align='center', wrap=True)
    ws.row_dimensions[4].height = 42
    for rr in range(5, 16):
        for c in range(3, 8):
            put(ws, (rr, c), None)
    put(ws, (5, 2), 'СОБСТВЕННЫЙ КАПИТАЛ', bold=True)
    put(ws, (6, 2), BORROWER, italic=True)
    ws.merge_cells(start_row=6, start_column=2, end_row=6, end_column=7)
    rows_eq = [('Фиксированные активы', r_sfa), ('Рабочий капитал', SH['wc']),
               ('Денежные средства (финансовые издержки льготного периода, расходы подготовительного периода, резерв)', None)]
    r = 7
    for lab, sr in rows_eq:
        put(ws, (r, 2), lab, wrap=True)
        if sr:
            put(ws, (r, 3), f"='Share of costs'!C{sr}", fmt=NUM); put(ws, (r, 4), f"='Share of costs'!D{sr}", fmt=NUM)
        else:
            put(ws, (r, 3), f"='Share of costs'!C{SH['fin']}+'Share of costs'!C{SH['res']}", fmt=NUM)
            put(ws, (r, 4), f"='Share of costs'!D{SH['fin']}+'Share of costs'!D{SH['res']}", fmt=NUM)
        put(ws, (r, 5), f'=C{r}+D{r}', fmt=NUM); put(ws, (r, 6), f'=IFERROR(E{r}/$E$10,0)', fmt=PCT)
        r += 1
    put(ws, (10, 2), 'ВСЕГО', bold=True, fill=FILL_TOT)
    for c in (3, 4, 5):
        put(ws, (10, c), f'=SUM({L(c)}7:{L(c)}9)', fmt=NUM, bold=True, fill=FILL_TOT)
    put(ws, (10, 6), '=IFERROR(E10/$E$10,0)', fmt=PCT, bold=True, fill=FILL_TOT)
    put(ws, (10, 7), '=IFERROR(E10/$E$15,0)', fmt=PCT, bold=True, fill=FILL_TOT)
    put(ws, (11, 2), 'КРЕДИТЫ', bold=True)
    put(ws, (12, 2), f'{BANK} (7% годовых, 84 мес., льготный период 12 мес.)', wrap=True)
    ws.row_dimensions[12].height = 42
    put(ws, (12, 3), '=C13', fmt=NUM); put(ws, (12, 4), '=D13', fmt=NUM); put(ws, (12, 5), '=C12+D12', fmt=NUM)
    put(ws, (12, 6), '=IFERROR(E12/$E$14,0)', fmt=PCT)
    put(ws, (13, 2), '   Фиксированные активы (30% аванс по EPC-контракту)')
    put(ws, (13, 3), f"='Share of costs'!E{r_sfa}", fmt=NUM); put(ws, (13, 4), f"='Share of costs'!F{r_sfa}", fmt=NUM)
    put(ws, (13, 5), '=C13+D13', fmt=NUM); put(ws, (13, 6), '=IFERROR(E13/$E$14,0)', fmt=PCT)
    put(ws, (14, 2), 'ВСЕГО', bold=True, fill=FILL_TOT)
    for c in (3, 4, 5):
        put(ws, (14, c), f'={L(c)}12', fmt=NUM, bold=True, fill=FILL_TOT)
    put(ws, (14, 6), '=IFERROR(E14/$E$14,0)', fmt=PCT, bold=True, fill=FILL_TOT)
    put(ws, (14, 7), '=IFERROR(E14/$E$15,0)', fmt=PCT, bold=True, fill=FILL_TOT)
    put(ws, (15, 2), 'ВСЕГО ФИНАНСИРОВАНИЯ', bold=True, fill=FILL_KEY)
    for c in (3, 4, 5):
        put(ws, (15, c), f'={L(c)}14+{L(c)}10', fmt=NUM, bold=True, fill=FILL_KEY)
    put(ws, (15, 7), '=IFERROR(E15/$E$15,0)', fmt=PCT, bold=True, fill=FILL_KEY)
    FP_EQ, FP_LOAN, FP_TOT = 'FinPlan!$E$10', 'FinPlan!$E$14', 'FinPlan!$E$15'

    # ------------------------------------------------------------------ ProdPlan (Приложение 5)
    ws = ws_pp
    title(ws, 'ПРОИЗВОДСТВЕННЫЙ ПЛАН', 'Приложение 5', last_col=11)
    for i, h in enumerate(['Н А И М Е Н О В А Н И Е', 'Ед. изм.', 'Мощность в год', 'Рынок сбыта',
                           'Цена за единицу, USD (при проектной мощности)', 'Выручка при полной мощности: экспорт', 'Выручка при полной мощности: внутр. рынок']):
        put(ws, (4, 2 + i), h, bold=True, fill=FILL_HEAD, align='center', wrap=True)
    ws.row_dimensions[4].height = 56
    r = 5
    PP = {}
    for i, p in enumerate(prods):
        put(ws, (r, 2), p['name_ru']); put(ws, (r, 3), p['unit'], align='center')
        put(ws, (r, 4), f'={src.ref(f"cap{i}")}', fmt=NUM1)
        put(ws, (r, 5), 'экспорт' if p['market'] == 'export' else 'внутренний рынок', align='center')
        put(ws, (r, 6), f'={X(f"price{i}", REF)}', fmt=NUM2)
        put(ws, (r, 7), f'=D{r}*F{r}' if p['market'] == 'export' else 0, fmt=NUM)
        put(ws, (r, 8), f'=D{r}*F{r}' if p['market'] != 'export' else 0, fmt=NUM)
        PP['reg' + str(i)] = r
        r += 1
    put(ws, (r, 2), 'Всего', bold=True, fill=FILL_TOT)
    for c in range(3, 9):
        put(ws, (r, c), f'=SUM({L(c)}5:{L(c)}{r - 1})' if c in (7, 8) else None, fmt=NUM, bold=True, fill=FILL_TOT)
    r += 2
    year_header(ws, r, total='Итого')
    r += 2
    def block(label):
        nonlocal r
        put(ws, (r, 2), label, bold=True, fill=FILL_TOT)
        for c in range(3, 11):
            put(ws, (r, c), None, fill=FILL_TOT)
        r += 1
    block('ОСВОЕНИЕ МОЩНОСТЕЙ')
    for i, p in enumerate(prods):
        PP[f'util{i}'] = yrow(ws, r, p['name_ru'], lambda y, c, i=i: f'={X(f"util{i}", y)}', unit='%', fmt=PCT, total='avg'); r += 1
    block('КОЛИЧЕСТВО ПРОИЗВОДСТВА')
    for i, p in enumerate(prods):
        PP[f'prod{i}'] = yrow(ws, r, p['name_ru'], lambda y, c, i=i: f'=$D${PP["reg" + str(i)]}*{c}{PP[f"util{i}"]}',
                              unit=p['unit'], fmt=NUM1); r += 1
    block('КОЛИЧЕСТВО ДЛЯ РЕАЛИЗАЦИИ')
    for i, p in enumerate(prods):
        PP[f'sexp{i}'] = yrow(ws, r, p['name_ru'] + ' — экспорт', lambda y, c, i=i: f'={X(f"sexp{i}", y)}', unit=p['unit'], fmt=NUM1); r += 1
        PP[f'sdom{i}'] = yrow(ws, r, p['name_ru'] + ' — внутренний рынок', lambda y, c, i=i: f'={X(f"sdom{i}", y)}', unit=p['unit'], fmt=NUM1); r += 1
    block('КОЛИЧЕСТВО ЗАПАСА ГОТОВОЙ ПРОДУКЦИИ НА КОНЕЦ ГОДА')
    for i, p in enumerate(prods):
        def fstock(y, c, i=i):
            prev = '0' if y == 2027 else f'{L(YC[y] - 1)}{r}'
            return f'={prev}+{c}{PP[f"prod{i}"]}-{c}{PP[f"sexp{i}"]}-{c}{PP[f"sdom{i}"]}'
        PP[f'stock{i}'] = yrow(ws, r, p['name_ru'], fstock, unit=p['unit'], fmt=NUM1, total='last'); r += 1
    block('ЦЕНА РЕАЛИЗАЦИИ')
    for i, p in enumerate(prods):
        PP[f'price{i}'] = yrow(ws, r, p['name_ru'], lambda y, c, i=i: f'={X(f"price{i}", y)}', unit=f'USD/{p["unit"]}', fmt=NUM2, total='avg'); r += 1
        ws.cell(row=PP['reg' + str(i)], column=6).value = f'={L(YC[REF])}{PP[f"price{i}"]}'
    if len(prods) > 1:
        r += 1
        put(ws, (r, 2), 'Попутная продукция (медный шлак, кальциевая селитра) в базовом сценарии ТЭО выручки не формирует — '
            'отражена в натуральном выражении (консервативный подход).', italic=True, border=False)

    # ------------------------------------------------------------------ SalePlan (Приложение 6)
    ws = ws_sp
    title(ws, 'ПЛАН ПРОДАЖ (в экв. долл. США, без НДС)', 'Приложение 6', last_col=11)
    year_header(ws, 4)
    r = 6
    SP = {}
    for sect, kk in [('ЭКСПОРТ', 'sexp'), ('ВНУТРЕННИЕ ПРОДАЖИ', 'sdom')]:
        put(ws, (r, 2), sect, bold=True, fill=FILL_TOT)
        for c in range(3, 11):
            put(ws, (r, c), None, fill=FILL_TOT)
        r += 1
        first_r = r
        for i, p in enumerate(prods):
            yrow(ws, r, p['name_ru'], lambda y, c, i=i, kk=kk: f'=ProdPlan!{c}{PP[f"{kk}{i}"]}*ProdPlan!{c}{PP[f"price{i}"]}', indent=1)
            r += 1
        SP[kk] = yrow(ws, r, 'Всего', lambda y, c, a=first_r, b=r - 1: f'=SUM({c}{a}:{c}{b})', bold=True)
        r += 2
    SP['tot'] = yrow(ws, r, 'ВСЕГО ПРОДАЖ', lambda y, c: f'={c}{SP["sexp"]}+{c}{SP["sdom"]}', bold=True, fill=FILL_KEY)
    r += 1
    yrow(ws, r, 'Доля экспорта в выручке', lambda y, c: f'=IFERROR({c}{SP["sexp"]}/{c}{SP["tot"]},0)', unit='%', fmt=PCT, total=None)

    # ------------------------------------------------------------------ RawMat (Приложение 7)
    ws = ws_rm
    title(ws, 'СТОИМОСТЬ СЫРЬЯ И МАТЕРИАЛОВ ПРИ ПРОЕКТНОЙ МОЩНОСТИ (в экв. долл. США)', 'Приложение 7', last_col=7,
          sub=f'Расчёт на {REFL} (освоение мощности {util[REF]:.0%}); цены в номинальном выражении; без НДС.')
    for i, h in enumerate(['НАИМЕНОВАНИЕ', 'Ед. изм.', 'Цена за единицу, USD', 'Годовая потребность', 'Годовая стоимость, USD', 'ДОЛЯ (%)', 'Примечание']):
        put(ws, (5, 2 + i), h, bold=True, fill=FILL_HEAD, align='center', wrap=True)
    ws.row_dimensions[5].height = 30
    r = 6
    mat_idx = [j for j, v in enumerate(D['varlines']) if v['element'] not in ENERGY]
    RM = {}
    sect_rows = []
    for sect, cur in [('СЫРЬЁ, МАТЕРИАЛЫ И ПЕРЕМЕННЫЕ ЗАТРАТЫ С ЦЕНОЙ В ТВЁРДОЙ ВАЛЮТЕ (USD)', 'USD'),
                      ('СЫРЬЁ, МАТЕРИАЛЫ И ПЕРЕМЕННЫЕ ЗАТРАТЫ С ЦЕНОЙ В НАЦИОНАЛЬНОЙ ВАЛЮТЕ (UZS)', 'UZS')]:
        put(ws, (r, 2), sect, bold=True, fill=FILL_TOT)
        for c in range(3, 9):
            put(ws, (r, c), None, fill=FILL_TOT)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        r += 1
        a = r
        for j in mat_idx:
            vl = D['varlines'][j]
            if vl['ccy'] != cur:
                continue
            put(ws, (r, 2), vl['name_ru'], wrap=True); put(ws, (r, 3), vl['unit'], align='center')
            put(ws, (r, 5), f'={X(f"vq{j}", REF)}', fmt=NUM1)
            put(ws, (r, 6), f'={X(f"vc{j}", REF)}', fmt=NUM)
            put(ws, (r, 4), f'=IFERROR(F{r}/E{r},0)', fmt=NUM2)
            put(ws, (r, 8), EL_RU[vl['element']], color='595959')
            r += 1
        for c in (3, 4, 5, 8):
            put(ws, (r, c), None)
        put(ws, (r, 2), 'ВСЕГО ' + ('В ТВЁРДОЙ ВАЛЮТЕ' if cur == 'USD' else 'В НАЦИОНАЛЬНОЙ ВАЛЮТЕ'), bold=True)
        put(ws, (r, 6), f'=SUM(F{a}:F{r - 1})' if r > a else 0, fmt=NUM, bold=True)
        sect_rows.append(r)
        RM[cur] = r
        r += 2
    RM['tot'] = r
    for c in (3, 4, 5, 8):
        put(ws, (r, c), None, fill=FILL_KEY)
    put(ws, (r, 2), 'ИТОГО СЫРЬЁ, МАТЕРИАЛЫ И ПРОЧИЕ ПЕРЕМЕННЫЕ ЗАТРАТЫ', bold=True, fill=FILL_KEY)
    put(ws, (r, 6), f'=F{RM["USD"]}+F{RM["UZS"]}', fmt=NUM, bold=True, fill=FILL_KEY)
    for rr in range(7, r + 1):
        if ws.cell(row=rr, column=6).value is not None:
            put(ws, (rr, 7), f'=IFERROR(F{rr}/$F${r},0)', fmt=PCT)
    r += 2
    put(ws, (r, 2), 'ПРОЧИЕ ПРОИЗВОДСТВЕННЫЕ РАСХОДЫ (постоянные)', bold=True, fill=FILL_TOT)
    for c in range(3, 9):
        put(ws, (r, c), None, fill=FILL_TOT)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    r += 1
    for c in range(2, 9):
        put(ws, (r, c), {2: 'Статьи затрат', 6: 'Годовая стоимость, USD', 8: 'Состав затрат'}.get(c), bold=True, fill=FILL_HEAD,
            align='center', wrap=True)
    ws.row_dimensions[r].height = 30
    r += 1
    for rr in range(r, r + 4):
        for c in (3, 4, 5, 7, 8):
            put(ws, (rr, c), None)
    a = r
    def var_sum(element, y):
        idx = [j for j, v in enumerate(D['varlines']) if v['element'] == element]
        return ('+'.join(X(f'vc{j}', y) for j in idx)) if idx else '0'
    fixnames = lambda el: '; '.join(f['name_ru'] for f in D['fixlines'] if f['element'] == el)[:250]
    RM['rep'] = r
    put(ws, (r, 2), 'Ремонт и техническое обслуживание'); put(ws, (r, 6), f'={X("el_repairs", REF)}-({var_sum("Repairs", REF)})', fmt=NUM)
    put(ws, (r, 8), fixnames('Repairs'), wrap=True, color='595959'); r += 1
    RM['oth'] = r
    put(ws, (r, 2), 'Прочие производственные расходы (охрана труда, накладные и др.)')
    put(ws, (r, 6), f'={X("el_other", REF)}-({var_sum("Other", REF)})-{X("met", REF)}', fmt=NUM)
    put(ws, (r, 8), fixnames('Other'), wrap=True, color='595959'); r += 1
    RM['met'] = r
    put(ws, (r, 2), 'Налог за пользование недрами (НДПИ)'); put(ws, (r, 6), f'={X("met", REF)}', fmt=NUM); r += 1
    put(ws, (r, 2), 'Итого:', bold=True); put(ws, (r, 6), f'=SUM(F{a}:F{r - 1})', fmt=NUM, bold=True); RM['other'] = r

    # ------------------------------------------------------------------ Labour (Приложение 8)
    ws = ws_lb
    title(ws, 'ЭНЕРГОРЕСУРСЫ И ОПЛАТА ТРУДА ПРИ ПРОЕКТНОЙ МОЩНОСТИ (в экв. долл. США)', 'Приложение 8', last_col=7,
          sub=f'Расчёт на {REFL}.')
    for i, h in enumerate(['СТАТЬИ ЗАТРАТ', 'Ед. изм.', 'Годовая потребность', 'Цена за единицу, USD', 'Годовая стоимость, USD']):
        put(ws, (5, 2 + i), h, bold=True, fill=FILL_HEAD, align='center', wrap=True)
    ws.row_dimensions[5].height = 30
    r = 6
    a = r
    for j, vl in enumerate(D['varlines']):
        if vl['element'] not in ENERGY:
            continue
        put(ws, (r, 2), vl['name_ru'], wrap=True); put(ws, (r, 3), vl['unit'], align='center')
        put(ws, (r, 4), f'={X(f"vq{j}", REF)}', fmt=NUM1); put(ws, (r, 6), f'={X(f"vc{j}", REF)}', fmt=NUM)
        put(ws, (r, 5), f'=IFERROR(F{r}/D{r},0)', fmt='#,##0.0000')
        r += 1
    en_var = (a, r - 1)
    for el, k in [('Electricity', 'el_electricity'), ('Gas / fuel', 'el_gas'), ('Water', 'el_water')]:
        put(ws, (r, 2), f'{EL_RU[el]} - постоянная часть')
        put(ws, (r, 6), f'={X(k, REF)}-({var_sum(el, REF)})', fmt=NUM); r += 1
    en_fix = (en_var[1] + 1, r - 1)
    for rr in range(en_fix[0], r + 1):
        for c in (3, 4, 5):
            put(ws, (rr, c), None, fill=FILL_TOT if rr == r else None)
    put(ws, (r, 2), 'ВСЕГО ЭНЕРГОРЕСУРСЫ', bold=True, fill=FILL_TOT)
    put(ws, (r, 6), f'=SUM(F{a}:F{r - 1})', fmt=NUM, bold=True, fill=FILL_TOT); LB = {'energy': r, 'en_var': en_var, 'en_fix': en_fix}
    r += 2
    put(ws, (r, 2), 'РАСЧЕТ ЗАРАБОТНОЙ ПЛАТЫ', bold=True, border=False)
    put(ws, (r, 6), 'Ставка социального налога', align='right', border=False)
    put(ws, (r, 7), f'={src.ref("soc_rate")}', fmt=PCT); r_soc = r; r += 1
    for i, h in enumerate(['ПОКАЗАТЕЛИ', 'Количество, чел.', 'Среднемесячная з/плата на 1 чел., сум',
                           'Среднемесячная з/плата на 1 чел., USD', 'Годовой фонд з/платы, USD', 'Отчисления (соц. налог), USD']):
        put(ws, (r, 2 + i), h, bold=True, fill=FILL_HEAD, align='center', wrap=True)
    ws.row_dimensions[r].height = 42
    r += 1
    for grp, is_ga in [('ПРОИЗВОДСТВЕННЫЙ ПЕРСОНАЛ', False), ('АДМИНИСТРАЦИЯ', True)]:
        put(ws, (r, 2), grp, bold=True, fill=FILL_TOT)
        for c in range(3, 8):
            put(ws, (r, c), None, fill=FILL_TOT)
        r += 1
        a = r
        for j, ps in enumerate(D['positions']):
            if (ps['cc'] == 'G&A') != is_ga:
                continue
            put(ws, (r, 2), ps['name_ru'], wrap=True)
            put(ws, (r, 3), f'={X(f"hc{j}", REF)}', fmt=NUM1)
            put(ws, (r, 4), f'={X(f"mc{j}", REF)}' if ps['ccy'] == 'UZS' else '-', fmt=NUM)
            put(ws, (r, 6), f'={X(f"pc{j}", REF)}', fmt=NUM)
            put(ws, (r, 5), f'=IFERROR(F{r}/C{r}/12,0)', fmt=NUM)
            put(ws, (r, 7), f'=F{r}*$G${r_soc}', fmt=NUM)
            r += 1
        put(ws, (r, 2), 'ВСЕГО', bold=True)
        for c in (4, 5):
            put(ws, (r, c), None)
        for c in (3, 6, 7):
            put(ws, (r, c), f'=SUM({L(c)}{a}:{L(c)}{r - 1})', fmt=NUM1 if c == 3 else NUM, bold=True)
        rr_tot = r
        r += 1
        for c in (3, 4, 5, 7):
            put(ws, (r, c), None, fill=FILL_TOT)
        put(ws, (r, 2), 'ИТОГО (фонд оплаты труда + социальный налог)', bold=True, fill=FILL_TOT)
        put(ws, (r, 6), f'=F{rr_tot}+G{rr_tot}', fmt=NUM, bold=True, fill=FILL_TOT)
        LB['ga' if is_ga else 'prod'] = r
        LB[('ga' if is_ga else 'prod') + '_hc'] = rr_tot
        r += 2
    for c in (4, 5, 7):
        put(ws, (r, c), None, fill=FILL_KEY)
    put(ws, (r, 2), 'ОБЩАЯ СУММА', bold=True, fill=FILL_KEY)
    put(ws, (r, 3), f'=C{LB["prod_hc"]}+C{LB["ga_hc"]}', fmt=NUM1, bold=True, fill=FILL_KEY)
    put(ws, (r, 6), f'=F{LB["prod"]}+F{LB["ga"]}', fmt=NUM, bold=True, fill=FILL_KEY); LB['all'] = r
    r += 1
    put(ws, (r, 2), 'Заработная плата установлена в сумах и пересчитана в USD с учётом индексации и курса.', italic=True, border=False)

    # ------------------------------------------------------------------ LOANS-dependent: CostSold (Приложение 11)
    ws = ws_cs
    title(ws, 'ГОДОВАЯ СТОИМОСТЬ ПРОДАННОЙ ПРОДУКЦИИ (в экв. долл. США)', 'Приложение 11', last_col=11)
    year_header(ws, 4)
    r = 6
    CS = {}
    CS['util'] = yrow(ws, r, 'Освоение мощностей (основная продукция)', lambda y, c: f'=ProdPlan!{c}{PP["util0"]}', unit='%', fmt=PCT, total='avg'); r += 2
    put(ws, (r, 2), 'ПРОИЗВОДСТВЕННАЯ СЕБЕСТОИМОСТЬ ПО ЭЛЕМЕНТАМ', bold=True, border=False); r += 1
    a = r
    for el, k in el_keys.items():
        if el == 'Depreciation':
            CS['dep'] = yrow(ws, r, 'Амортизация (Приложение 4)', lambda y, c: f'={DEPR("cos", y)}', indent=1); r += 1
            continue
        if el == 'Other':
            CS['other'] = yrow(ws, r, EL_RU[el] + ' (без НДПИ)', lambda y, c: f'={X(k, y)}-{X("met", y)}', indent=1); r += 1
            CS['met'] = yrow(ws, r, 'Налог за пользование недрами (НДПИ)', lambda y, c: f'={X("met", y)}', indent=1); r += 1
            continue
        CS[k] = yrow(ws, r, EL_RU[el], lambda y, c, k=k: f'={X(k, y)}', indent=1); r += 1
    CS['prodcost'] = yrow(ws, r, 'Себестоимость произведенной продукции', lambda y, c: f'=SUM({c}{a}:{c}{r - 1})', bold=True, fill=FILL_TOT); r += 1
    CS['stock'] = yrow(ws, r, 'Изменение остатков готовой продукции (− прирост запаса)',
                       lambda y, c: f'={X("cos", y)}-{X("el_total", y)}-{X("capbc_dep", y)}', indent=1,
                       note='Разница между себестоимостью продаж и производственной себестоимостью по ТЭО (запасы готовой продукции).'); r += 1
    CS['cos'] = yrow(ws, r, 'Себестоимость реализованной продукции', lambda y, c: f'={c}{CS["prodcost"]}+{c}{CS["stock"]}', bold=True, fill=FILL_TOT); r += 1
    cash_el = [k for el, k in el_keys.items() if el != 'Depreciation']
    CS['usd_sh'] = yrow(ws, r, 'Доля затрат в твёрдой валюте (USD) в денежной себестоимости',
                        lambda y, c: f'=MIN(1,IFERROR({X("usd_prod", y)}/({"+".join(X(k, y) for k in cash_el)}),0))', unit='%', fmt=PCT, indent=1); r += 1
    CS['cos_usd'] = yrow(ws, r, 'в т.ч. затраты в твёрдой валюте (USD)', lambda y, c: f'={c}{CS["cos"]}*{c}{CS["usd_sh"]}', indent=1); r += 1
    CS['cos_uzs'] = yrow(ws, r, 'в т.ч. затраты в национальной валюте', lambda y, c: f'={c}{CS["cos"]}-{c}{CS["cos_usd"]}', indent=1); r += 2
    put(ws, (r, 2), 'РАСХОДЫ ПЕРИОДА', bold=True, border=False); r += 1
    CS['ga_pers'] = yrow(ws, r, 'Зарплата административного персонала (с соц. налогом)', lambda y, c: f'={X("pers_ga", y)}', indent=1); r += 1
    CS['ga_da'] = yrow(ws, r, 'Амортизация административных активов', lambda y, c: f'={DEPR("ga", y)}', indent=1); r += 1
    CS['ga_oth'] = yrow(ws, r, 'Прочие административные расходы (аренда, налог за воду, прочие)',
                        lambda y, c: f'={X("ga", y)}-{X("pers_ga", y)}-{X("opex_ga_da", y)}', indent=1); r += 1
    CS['sell'] = yrow(ws, r, 'Коммерческие расходы (расходы по реализации)', lambda y, c: f'={X("sell", y)}', indent=1); r += 1
    CS['period'] = yrow(ws, r, 'Итого расходы периода', lambda y, c: f'=SUM({c}{CS["ga_pers"]}:{c}{CS["sell"]})', bold=True, fill=FILL_TOT); r += 2
    put(ws, (r, 2), 'НАЛОГИ И ПРОЧИЕ ОПЕРАЦИОННЫЕ РАСХОДЫ', bold=True, border=False); r += 1
    CS['tx_prop'] = yrow(ws, r, 'Налог на имущество', lambda y, c: f'={X("prop_tax", y)}', indent=1); r += 1
    CS['tx_land'] = yrow(ws, r, 'Земельный налог', lambda y, c: f'={X("land_tax", y)}', indent=1); r += 1
    CS['tx_rent'] = yrow(ws, r, 'Специальный рентный налог', lambda y, c: f'={X("rent_tax", y)}', indent=1); r += 1
    CS['cont'] = yrow(ws, r, 'Резерв на непредвиденные операционные расходы', lambda y, c: f'={X("contingency", y)}', indent=1); r += 1
    CS['ooe_oth'] = yrow(ws, r, 'Прочие операционные расходы / (доходы)',
                         lambda y, c: f'={X("ooe", y)}-{X("ooi", y)}-{X("prop_tax", y)}-{X("land_tax", y)}-{X("rent_tax", y)}-{X("contingency", y)}', indent=1); r += 1
    ws_rm.cell(row=RM['met'], column=6).value = f'=CostSold!{L(YC[REF])}{CS["met"]}'
    CS['opx'] = yrow(ws, r, 'Итого налоги и прочие операционные расходы', lambda y, c: f'=SUM({c}{CS["tx_prop"]}:{c}{CS["ooe_oth"]})', bold=True, fill=FILL_TOT); r += 2
    put(ws, (r, 2), 'ФИНАНСОВЫЕ ИЗДЕРЖКИ', bold=True, border=False); r += 1
    CS['int'] = yrow(ws, r, 'Проценты за кредит (1-й год - капитализированы, оплачены за счёт собственных средств)',
                     lambda y, c: '=0' if y == 2027 else f'={LNR("int", y)}', indent=1); r += 2
    CS['full'] = yrow(ws, r, 'ПОЛНАЯ СТОИМОСТЬ ПРОДАННОЙ ПРОДУКЦИИ', lambda y, c: f'={c}{CS["cos"]}+{c}{CS["period"]}+{c}{CS["opx"]}+{c}{CS["int"]}',
                      bold=True, fill=FILL_KEY); r += 1
    CS['unit'] = yrow(ws, r, f'Полная себестоимость единицы основной продукции', lambda y, c: f'=IFERROR({c}{CS["full"]}/(ProdPlan!{c}{PP["sexp0"]}+ProdPlan!{c}{PP["sdom0"]}),0)',
                      unit=f'USD/{main["unit"]}', fmt=NUM, total=None); r += 1

    # ------------------------------------------------------------------ Taxes (Приложение 9) — CIT needs PBT from ProfLoss
    # ProfLoss rows are fixed below; define row numbers first
    PLR = {}
    rr = 6
    for k in ['rev_g', 'rev_exp', 'rev_dom', 'vat', 'rev', 'cos', 'cos_usd', 'cos_uzs', 'gp', 'period', 'opx', 'ebit',
              'fin', 'int', 'pbt', 'cit', 'np']:
        PLR[k] = rr
        rr += 1

    ws = ws_tx
    title(ws, 'НАЛОГИ И СБОРЫ (в экв. долл. США)', 'Приложение 9', last_col=11)
    year_header(ws, 4, unit='Ставка')
    r = 6
    TX = {}
    put(ws, (r, 2), 'Налоги и сборы, включаемые в расходы', bold=True, fill=FILL_TOT)
    TX['inc'] = r; r += 1
    a = r
    TX['prop'] = yrow(ws, r, 'Налог на имущество', lambda y, c: f'=CostSold!{c}{CS["tx_prop"]}', unit=f'{meta["prop_rate_a"]:.2%}' if meta.get('prop_rate_a') else '', indent=1); r += 1
    TX['land'] = yrow(ws, r, 'Земельный налог', lambda y, c: f'=CostSold!{c}{CS["tx_land"]}', unit='', indent=1); r += 1
    TX['rent'] = yrow(ws, r, 'Специальный рентный налог', lambda y, c: f'=CostSold!{c}{CS["tx_rent"]}', unit=f'{meta["rent_rate"]:.0%}' if meta['rent_rate'] else '', indent=1); r += 1
    TX['met'] = yrow(ws, r, 'Налог за пользование недрами (НДПИ) — в себестоимости', lambda y, c: f'=CostSold!{c}{CS["met"]}', unit='', indent=1); r += 1
    TX['soc'] = yrow(ws, r, 'Социальный налог (производственный и административный персонал)',
                     lambda y, c: f'={X("el_social", y)}+{X("pers_ga", y)}-({"+".join(X(f"pc{j}", y) for j, ps in enumerate(D["positions"]) if ps["cc"] == "G&A") or "0"})',
                     unit=f'{meta["soc_rate"]:.0%}', indent=1); r += 1
    for c in range(4, 11):
        put(ws, (TX['inc'], c), f'=SUM({L(c)}{a}:{L(c)}{r - 1})', fmt=NUM, bold=True, fill=FILL_TOT)
    put(ws, (TX['inc'], 3), None, fill=FILL_TOT)
    r += 1
    put(ws, (r, 2), 'Прочие налоги', bold=True, fill=FILL_TOT); TX['oth'] = r; r += 1
    TX['cit'] = yrow(ws, r, 'Налог на прибыль (Приложение 13)', lambda y, c: f'=ProfLoss!{c}{PLR["cit"]}', unit='', indent=1); r += 1
    r_vat = r
    r_vo, r_vio, r_vic, r_vnet = r_vat + 4, r_vat + 5, r_vat + 6, r_vat + 7
    TX['vat'] = yrow(ws, r, 'НДС к уплате в бюджет (при превышении входящего НДС - к возмещению, см. ниже)',
                     lambda y, c: f'=MAX(0,{c}{r_vnet})', unit='12%', indent=1); r += 1
    for c in range(4, 11):
        put(ws, (TX['oth'], c), f'={L(c)}{TX["cit"]}+{L(c)}{TX["vat"]}', fmt=NUM, bold=True, fill=FILL_TOT)
    put(ws, (TX['oth'], 3), None, fill=FILL_TOT)
    TX['all'] = yrow(ws, r, 'ВСЕГО НАЛОГОВ И ОТЧИСЛЕНИЙ', lambda y, c: f'={c}{TX["inc"]}+{c}{TX["oth"]}', unit='', bold=True, fill=FILL_KEY); r += 2
    put(ws, (r, 2), 'Расчет НДС', bold=True, border=False); r += 1
    assert r == r_vo
    TX['vo'] = yrow(ws, r_vo, 'НДС исходящий (реализация на внутреннем рынке; экспорт — 0%)', lambda y, c: f'={X("vat_out", y)}', unit='')
    TX['vio'] = yrow(ws, r_vio, 'НДС входящий — по операционным затратам', lambda y, c: f'={X("vat_in_op", y)}', unit='')
    TX['vic'] = yrow(ws, r_vic, 'НДС входящий - по капитальным вложениям (1-й год вкл. подготовительный период)',
                     lambda y, c: f'={X("vat_in_capex", y)}' + (f'+{X("vat_in_capex", 2026)}' if y == 2027 else ''), unit='')
    TX['vnet'] = yrow(ws, r_vnet, 'НДС к уплате (+) / к возмещению из бюджета (−)', lambda y, c: f'={c}{r_vo}-{c}{r_vio}-{c}{r_vic}', unit='', bold=True,
                      note='Для экспортных проектов (ставка 0%) входящий НДС подлежит возмещению; движение НДС учтено в оборотном капитале (Приложение 14).')
    r = r_vnet + 2
    put(ws, (r, 2), 'Расчет налога на прибыль', bold=True, border=False); r += 1
    TX['pbt'] = yrow(ws, r, 'Прибыль (убыток) до налогообложения (Приложение 13)', lambda y, c: f'=ProfLoss!{c}{PLR["pbt"]}', unit=''); r += 1
    TX['adj'] = yrow(ws, r, 'Налоговые корректировки (вменённый доход по реализации со скидкой и др.)' if meta.get('rent_rate') else 'Налоговые корректировки', lambda y, c: f'={X("tax_adj", y)}', unit=''); r += 1
    TX['base'] = yrow(ws, r, 'Налогооблагаемая прибыль', lambda y, c: f'={c}{TX["pbt"]}+{c}{TX["adj"]}', unit=''); r += 1
    TX['rate'] = yrow(ws, r, 'Ставка налога на прибыль', lambda y, c: f'={X("cit_rate", y)}', unit='%', fmt=PCT); r += 1
    TX['cit_bp'] = yrow(ws, r, 'Налог на прибыль', lambda y, c: f'={c}{TX["base"]}*{c}{TX["rate"]}', unit='', bold=True); r += 2

    # ------------------------------------------------------------------ ProfLoss (Приложение 13)
    ws = ws_pl
    title(ws, 'ПРЕДПОЛАГАЕМЫЕ СЧЕТА ПРИБЫЛЕЙ И УБЫТКОВ (в экв. долл. США)', 'Приложение 13', last_col=11)
    year_header(ws, 4)
    fx = {}
    def P(k, label, f, **kw):
        assert ws.max_row < PLR[k] or True
        yrow(ws, PLR[k], label, f, **kw)
    P('rev_g', 'Выручка от реализации продукции (с НДС)', lambda y, c: f'={c}{PLR["rev_exp"]}+{c}{PLR["rev_dom"]}', bold=True)
    P('rev_exp', 'экспорт (в твёрдой валюте)', lambda y, c: f'=SalePlan!{c}{SP["sexp"]}', indent=1)
    P('rev_dom', 'внутренний рынок (с НДС)', lambda y, c: f'=SalePlan!{c}{SP["sdom"]}+Taxes!{c}{TX["vo"]}', indent=1)
    P('vat', 'НДС, начисленный с реализации', lambda y, c: f'=Taxes!{c}{TX["vo"]}')
    P('rev', 'Чистая выручка от реализации продукции', lambda y, c: f'={c}{PLR["rev_g"]}-{c}{PLR["vat"]}', bold=True, fill=FILL_TOT)
    P('cos', 'Себестоимость реализованной продукции', lambda y, c: f'=CostSold!{c}{CS["cos"]}')
    P('cos_usd', 'в твёрдой валюте', lambda y, c: f'=CostSold!{c}{CS["cos_usd"]}', indent=1)
    P('cos_uzs', 'в местной валюте', lambda y, c: f'=CostSold!{c}{CS["cos_uzs"]}', indent=1)
    P('gp', 'Валовая прибыль (убыток)', lambda y, c: f'={c}{PLR["rev"]}-{c}{PLR["cos"]}', bold=True, fill=FILL_TOT)
    P('period', 'Расходы периода (административные и коммерческие)', lambda y, c: f'=CostSold!{c}{CS["period"]}')
    P('opx', 'Налоги и прочие операционные расходы', lambda y, c: f'=CostSold!{c}{CS["opx"]}')
    P('ebit', 'Прибыль (убыток) от основной деятельности (EBIT)', lambda y, c: f'={c}{PLR["gp"]}-{c}{PLR["period"]}-{c}{PLR["opx"]}', bold=True, fill=FILL_TOT)
    P('fin', 'Финансовые расходы', lambda y, c: f'={c}{PLR["int"]}')
    P('int', 'Проценты за кредит', lambda y, c: f'=CostSold!{c}{CS["int"]}', indent=1)
    P('pbt', 'Прибыль (убыток) до налогообложения', lambda y, c: f'={c}{PLR["ebit"]}-{c}{PLR["fin"]}', bold=True, fill=FILL_TOT)
    P('cit', 'Налог на прибыль (Приложение 9)', lambda y, c: f'=Taxes!{c}{TX["cit_bp"]}')
    P('np', 'Чистая прибыль (убыток)', lambda y, c: f'={c}{PLR["pbt"]}-{c}{PLR["cit"]}', bold=True, fill=FILL_KEY)
    r = PLR['np'] + 2
    put(ws, (r, 2), 'Коэффициенты рентабельности и покрытия', bold=True, border=False); r += 1
    PLR['gm'] = yrow(ws, r, 'Валовая рентабельность', lambda y, c: f'=IFERROR({c}{PLR["gp"]}/{c}{PLR["rev"]},0)', unit='%', fmt=PCT, total=None); r += 1
    PLR['nm'] = yrow(ws, r, 'Чистая рентабельность', lambda y, c: f'=IFERROR({c}{PLR["np"]}/{c}{PLR["rev"]},0)', unit='%', fmt=PCT, total=None); r += 1
    PLR['ebitda'] = yrow(ws, r, 'EBITDA', lambda y, c: f'={c}{PLR["ebit"]}+{DEPR("tot", y)}'); r += 1
    PLR['ebm'] = yrow(ws, r, 'Маржа EBITDA', lambda y, c: f'=IFERROR({c}{PLR["ebitda"]}/{c}{PLR["rev"]},0)', unit='%', fmt=PCT, total=None); r += 1
    PLR['icr'] = yrow(ws, r, 'EBIT / Проценты за кредит', lambda y, c: f'=IF({c}{PLR["int"]}>0,{c}{PLR["ebit"]}/{c}{PLR["int"]},"-")', unit='x', fmt=MULT, total=None); r += 1
    PLR['debt'] = yrow(ws, r, 'Задолженность по кредиту на конец года', lambda y, c: f'={LNR("close", y)}', total=None); r += 1
    PLR['de'] = yrow(ws, r, 'Долг / EBITDA', lambda y, c: f'=IFERROR({c}{PLR["debt"]}/{c}{PLR["ebitda"]},0)', unit='x', fmt=MULT, total=None); r += 2
    put(ws, (r, 2), '1-й год: проценты за кредит льготного периода оплачиваются за счёт собственных средств и капитализируются '
        'в стоимость основных средств (МСФО IAS 23, Приложения 1 и 4); в отчёт о прибылях и убытках не включаются.', italic=True, border=False)

    # ------------------------------------------------------------------ CostTotal (Приложение 10)
    ws = ws_ct
    title(ws, 'СТОИМОСТЬ ПРОДУКЦИИ ПРИ ПРОЕКТНОЙ МОЩНОСТИ (в экв. долл. США)', 'Приложение 10', last_col=8,
          sub=f'Расчёт на {REFL}; разделение на постоянные и переменные затраты.')
    for i, h in enumerate(['Статьи издержек', 'Всего стоимость', 'Доля постоянных затрат (%)', 'Доля переменных затрат (%)',
                           'Фиксированная стоимость', 'Переменная стоимость']):
        put(ws, (5, 2 + i), h, bold=True, fill=FILL_HEAD, align='center', wrap=True)
    ws.row_dimensions[5].height = 42
    R = REF
    mat_var_usd = '+'.join(X(f'vc{j}', R) for j in mat_idx if D['varlines'][j]['ccy'] == 'USD') or '0'
    mat_var_uzs = '+'.join(X(f'vc{j}', R) for j in mat_idx if D['varlines'][j]['ccy'] != 'USD') or '0'
    en_var = '+'.join(X(f'vc{j}', R) for j, v in enumerate(D['varlines']) if v['element'] in ENERGY) or '0'
    oth_var = var_sum('Other', R)
    ga_pos = '+'.join(X(f'pc{j}', R) for j, ps in enumerate(D['positions']) if ps['cc'] == 'G&A') or '0'
    cc = L(YC[R])
    ev, ef = LB['en_var'], LB['en_fix']
    items = [
        ('Сырьё, материалы и прочие переменные затраты', None, None, True),
        ('   с ценой в твёрдой валюте (USD)', '0', f'=RawMat!F{RM["USD"]}', False),
        ('   с ценой в национальной валюте', '0', f'=RawMat!F{RM["UZS"]}', False),
        ('Энергоресурсы (электроэнергия, газ, вода)', f'=SUM(Labour!F{ef[0]}:F{ef[1]})' if ef[1] >= ef[0] else '0',
         f'=SUM(Labour!F{ev[0]}:F{ev[1]})' if ev[1] >= ev[0] else '0', False),
        ('Зарплата производственного персонала (с соц. налогом)', f'=Labour!F{LB["prod"]}', '0', False),
        ('Зарплата административного персонала (с соц. налогом)', f'=Labour!F{LB["ga"]}', '0', False),
        ('Ремонт и техническое обслуживание (постоянная часть)', f'=RawMat!F{RM["rep"]}', '0', False),
        ('Прочие производственные расходы (постоянная часть)', f'=RawMat!F{RM["oth"]}', '0', False),
        ('Налог за пользование недрами (НДПИ)', '0', f'=RawMat!F{RM["met"]}', False),
        ('Административные расходы (без персонала и амортизации)', f'={X("opex_fixed", R)}',
         f'={X("opex_var", R)}-{X("sell", R)}', False),
        ('Коммерческие расходы', '0', f'=CostSold!{cc}{CS["sell"]}', False),
        ('Налог на имущество и земельный налог', f'=CostSold!{cc}{CS["tx_prop"]}+CostSold!{cc}{CS["tx_land"]}', '0', False),
        ('Специальный рентный налог', '0', f'=CostSold!{cc}{CS["tx_rent"]}', False),
        ('Резерв на непредвиденные операционные расходы', '0', f'=CostSold!{cc}{CS["cont"]}', False),
        ('Прочие операционные расходы / (доходы)', f'=CostSold!{cc}{CS["ooe_oth"]}', '0', False),
        ('Амортизация (Приложение 4)', f'=Depreciate!{L(7 + YEARS.index(R))}{DEP["tot"]}', '0', False),
    ]
    # materials: fixed residual of raw+aux elements, if any, goes to the hard/local split proportionally -> keep explicit
    r = 6
    CT = {}
    for lab, ffix, fvar, head in items:
        put(ws, (r, 2), lab, bold=head)
        if head:
            CT['mat'] = r
            r += 1
            continue
        put(ws, (r, 6), ffix, fmt=NUM); put(ws, (r, 7), fvar, fmt=NUM)
        put(ws, (r, 3), f'=F{r}+G{r}', fmt=NUM)
        put(ws, (r, 4), f'=IFERROR(F{r}/C{r},0)', fmt=PCT); put(ws, (r, 5), f'=IFERROR(G{r}/C{r},0)', fmt=PCT)
        r += 1
    # header row for materials = sum of the two sub-rows
    m = CT['mat']
    for c in (3, 6, 7):
        put(ws, (m, c), f'={L(c)}{m + 1}+{L(c)}{m + 2}', fmt=NUM, bold=True)
    put(ws, (m, 4), f'=IFERROR(F{m}/C{m},0)', fmt=PCT, bold=True); put(ws, (m, 5), f'=IFERROR(G{m}/C{m},0)', fmt=PCT, bold=True)
    # materials fixed residual (raw + aux elements minus their variable lines) -> add to row "в местной валюте" as fixed
    put(ws, (m + 2, 6), f'=({X("el_raw", R)}+{X("el_auxiliary", R)})-({var_sum("Raw materials", R)})-({var_sum("Auxiliary materials", R)})', fmt=NUM)
    CT['tot'] = r
    put(ws, (r, 2), 'ИТОГО', bold=True, fill=FILL_KEY)
    for c in (3, 6, 7):
        put(ws, (r, c), f'=SUM({L(c)}{m}:{L(c)}{r - 1})-{L(c)}{m}', fmt=NUM, bold=True, fill=FILL_KEY)
    put(ws, (r, 4), f'=IFERROR(F{r}/C{r},0)', fmt=PCT, bold=True, fill=FILL_KEY)
    put(ws, (r, 5), f'=IFERROR(G{r}/C{r},0)', fmt=PCT, bold=True, fill=FILL_KEY)
    r += 1
    put(ws, (r, 2), 'в т.ч. в местной валюте'); put(ws, (r, 3), f'=C{CT["tot"]}-C{r + 1}', fmt=NUM); r += 1
    put(ws, (r, 2), 'в т.ч. в твёрдой валюте (производственные затраты в USD)')
    put(ws, (r, 3), f'={X("usd_prod", R)}', fmt=NUM); r += 2
    put(ws, (r, 2), 'Выручка при проектной мощности'); put(ws, (r, 3), f'=SalePlan!{L(YC[R])}{SP["tot"]}', fmt=NUM); CT['rev'] = r; r += 1
    put(ws, (r, 2), 'Затраты на 1 USD выручки'); put(ws, (r, 3), f'=IFERROR(C{CT["tot"]}/C{CT["rev"]},0)', fmt='0.000'); r += 1
    put(ws, (r, 2), f'Полная себестоимость единицы основной продукции, USD/{main["unit"]}')
    put(ws, (r, 3), f'=IFERROR(C{CT["tot"]}/ProdPlan!{L(YC[R])}{PP["prod0"]},0)', fmt=NUM); r += 1

    # ------------------------------------------------------------------ BreakPoint (Приложение 16)
    ws = ws_bp
    title(ws, 'ТОЧКА БЕЗУБЫТОЧНОСТИ (в экв. долл. США)', 'Приложение 16', last_col=4, sub=f'Расчёт на {REFL} (проектная мощность).')
    rows_bp = [('Общие продажи при проектной мощности', f'=CostTotal!C{CT["rev"]}', NUM),
               ('в местной валюте', f'=SalePlan!{L(YC[R])}{SP["sdom"]}', NUM),
               ('в твердой валюте', f'=SalePlan!{L(YC[R])}{SP["sexp"]}', NUM),
               ('Фиксированная стоимость', f'=CostTotal!F{CT["tot"]}', NUM),
               ('Переменная стоимость', f'=CostTotal!G{CT["tot"]}', NUM),
               ('Маржинальный доход', '=C4-C8', NUM),
               ('Точка безубыточности (доля от выручки при проектной мощности)', '=IFERROR(C7/C9,0)', PCT),
               ('Точка безубыточности, USD выручки', '=C10*C4', NUM),
               ('Точка безубыточности в натуральном выражении (основная продукция)',
                f'=C10*ProdPlan!{L(YC[R])}{PP["prod0"]}', NUM1),
               ('Запас финансовой прочности', '=1-C10', PCT)]
    for i, (lab, f, fm) in enumerate(rows_bp):
        put(ws, (4 + i, 2), lab, bold=i in (0, 6)); put(ws, (4 + i, 3), f, fmt=fm, bold=i in (0, 6))
    put(ws, (4 + 8, 4), main['unit'], border=False)

    # ------------------------------------------------------------------ WorkCap (Приложение 14)
    ws = ws_wc
    title(ws, 'ПРЕДПОЛОЖЕНИЯ ПО РАБОЧЕМУ КАПИТАЛУ (в экв. долл. США)', 'Приложение 14', last_col=11)
    put(ws, (4, 2), 'Показатели', bold=True, fill=FILL_HEAD); put(ws, (4, 3), 'Дни покрытия', bold=True, fill=FILL_HEAD)
    params = [('Период поступления от продаж (дебиторская задолженность)', 'dso'), ('Запасы сырья и материалов на складе', 'rm_days'),
              ('Незавершенная продукция', 'wip_days'), ('Запасы готовой продукции (минимальный запас)', 'fg_days'),
              ('Период оплаты поставщикам (кредиторская задолженность)', 'dpo'), ('Задолженность по заработной плате', 'wage_lag'),
              ('Возмещение НДС (операционный)', 'vat_refund_lag'), ('Возмещение НДС по капитальным вложениям', 'capex_vat_lag')]
    r = 5
    for lab, k in params:
        put(ws, (r, 2), lab); put(ws, (r, 3), f'={src.ref(k)}', fmt='0'); r += 1
    r += 1
    rh = r
    year_header(ws, rh, label='ПОКАЗАТЕЛИ (на конец года)', unit='На начало строительства')
    r = rh + 2
    WC = {}
    def wrow(k, lab, f, bold=False, fill=None, opening=0):
        nonlocal r
        put(ws, (r, 2), lab, bold=bold, fill=fill)
        put(ws, (r, 3), opening, fmt=NUM, bold=bold, fill=fill)
        for y in YEARS:
            put(ws, (r, YC[y]), f(y, L(YC[y])), fmt=NUM, bold=bold, fill=fill)
        WC[k] = r
        r += 1
    wrow('rec', 'Дебиторская задолженность', lambda y, c: f'={X("bs_rec", y)}')
    wrow('inv', 'Запасы (сырьё, материалы, НЗП, готовая продукция)', lambda y, c: f'={X("bs_inv", y)}')
    wrow('vat', 'НДС к возмещению', lambda y, c: f'={X("bs_vat", y)}')
    wrow('ca', 'Итого текущих активов', lambda y, c: f'=SUM({c}{WC["rec"]}:{c}{WC["vat"]})', bold=True, fill=FILL_TOT)
    wrow('pay', 'Кредиторская задолженность', lambda y, c: f'={X("bs_pay", y)}')
    wrow('wag', 'Задолженность по оплате труда и налогам с ФОТ', lambda y, c: f'={X("bs_wag", y)}')
    wrow('cl', 'Итого текущих обязательств', lambda y, c: f'={c}{WC["pay"]}+{c}{WC["wag"]}', bold=True, fill=FILL_TOT)
    wrow('nwc', 'Чистый рабочий капитал', lambda y, c: f'={c}{WC["ca"]}-{c}{WC["cl"]}', bold=True, fill=FILL_KEY)
    r += 1
    def nwc_src(y):
        return f'({X("bs_rec", y)}+{X("bs_inv", y)}+{X("bs_vat", y)}-{X("bs_pay", y)}-{X("bs_wag", y)})'
    def prevc(y):
        return 'C' if y == 2027 else L(YC[y] - 1)
    wrow('dnwc', 'Изменения в рабочем капитале (− рост / + снижение)', lambda y, c: f'=-({c}{WC["nwc"]}-{prevc(y)}{WC["nwc"]})', bold=True)
    wrow('doth', 'Прочие изменения операционных статей (налоги к уплате и пр.)',
         lambda y, c: f'=({X("cgo", y)}-{X("ebitda", y)}' + (f'+{X("cgo", 2026)}-{X("ebitda", 2026)}' if y == 2027 else '') + f')+{nwc_src(y)}'
                      + (f'-{nwc_src(y - 1)}' if y > 2027 else ''))
    wrow('dwc', 'Изменение рабочего капитала для потока наличности (Приложение 15)',
         lambda y, c: f'={c}{WC["dnwc"]}+{c}{WC["doth"]}', bold=True, fill=FILL_TOT)
    for k in ('dnwc', 'doth', 'dwc'):
        ws.cell(row=WC[k], column=3).value = None
    ws_pc.cell(row=r_wc, column=4).value = f'=WorkCap!D{WC["rec"]}'
    ws_pc.cell(row=r_wc, column=3).value = f'=WorkCap!D{WC["nwc"]}-D{r_wc}'
    for c in (3, 4):
        ws_pc.cell(row=r_wc, column=c).number_format = NUM
        ws_pc.cell(row=r_wc, column=c).font = font(True)
        ws_pc.cell(row=r_wc, column=c).fill = FILL_TOT
        ws_pc.cell(row=r_wc, column=c).border = BORDER
    r += 1
    put(ws, (r, 2), '1-й год: изменение рассчитано от начала строительства, т.е. включает подготовительный период.',
        italic=True, border=False)

    # ------------------------------------------------------------------ CashFlow (Приложение 15)
    ws = ws_cf
    title(ws, 'ПОТОКИ НАЛИЧНОСТИ (в экв. долл. США)', 'Приложение 15', last_col=11,
          sub='1-й год - инвестиционный и льготный период; включает также освоение капвложений и расходы подготовительного периода.')
    year_header(ws, 4)
    r = 6
    CF = {}
    def C_(k, label, f, **kw):
        nonlocal r
        CF[k] = yrow(ws, r, label, f, **kw)
        r += 1
    C_('rev', 'Выручка от реализации продукции (с НДС)', lambda y, c: f'=ProfLoss!{c}{PLR["rev_g"]}')
    C_('vat', 'НДС, начисленный с реализации', lambda y, c: f'=-ProfLoss!{c}{PLR["vat"]}')
    C_('nrev', 'Чистая выручка от реализации продукции', lambda y, c: f'={c}{CF["rev"]}+{c}{CF["vat"]}', bold=True)
    C_('dwc', 'Изменения в рабочем капитале (Приложение 14)', lambda y, c: f'=WorkCap!{c}{WC["dwc"]}')
    C_('sales', 'Наличность от продаж', lambda y, c: f'={c}{CF["nrev"]}+{c}{CF["dwc"]}', bold=True, fill=FILL_TOT)
    C_('cos', 'Себестоимость реализованной продукции', lambda y, c: f'=-ProfLoss!{c}{PLR["cos"]}')
    C_('gross', 'Валовые поступления наличности', lambda y, c: f'={c}{CF["sales"]}+{c}{CF["cos"]}', bold=True, fill=FILL_TOT)
    C_('period', 'Расходы периода', lambda y, c: f'=-ProfLoss!{c}{PLR["period"]}')
    C_('tax', 'Налоги и прочие операционные расходы', lambda y, c: f'=-ProfLoss!{c}{PLR["opx"]}')
    C_('pre', 'Расходы подготовительного периода', lambda y, c: f'=-ProjCost!$E${PCR["pre"]}' if y == 2027 else '=0')
    C_('cit', 'Налог на прибыль', lambda y, c: f'=-ProfLoss!{c}{PLR["cit"]}')
    C_('da', 'Амортизация (неденежная статья, прибавляется)', lambda y, c: f'={DEPR("tot", y)}')
    C_('A', 'Операционный поток наличности (A)', lambda y, c: f'=SUM({c}{CF["gross"]}:{c}{CF["da"]})', bold=True, fill=FILL_KEY)
    C_('eq', 'Собственный капитал (взнос заемщика)', lambda y, c: f'={FP_EQ}' if y == 2027 else '=0')
    C_('loan', f'Привлеченный кредит {BANK_S}', lambda y, c: f'={LNR("drawn", y)}')
    C_('capex', 'Инвестиции (капитальные вложения, Приложение 4)', lambda y, c: f'=-{DEPR("capex", y)}')
    C_('B', 'Проценты за кредит (B)', lambda y, c: f'=-ProfLoss!{c}{PLR["int"]}')
    C_('fin1', 'Прочие финансовые издержки в инвестиционный период (проценты льготного периода)',
       lambda y, c: f'=-{I_GRACE}' if y == 2027 else '=0')
    C_('net1', 'Чистые поступления наличности', lambda y, c: f'={c}{CF["A"]}+{c}{CF["eq"]}+{c}{CF["loan"]}+{c}{CF["capex"]}+{c}{CF["B"]}+{c}{CF["fin1"]}',
       bold=True, fill=FILL_TOT)
    C_('C', 'Выплата основного долга (C)', lambda y, c: f'=-{LNR("princ", y)}')
    C_('net', 'Чистый поток наличности', lambda y, c: f'={c}{CF["net1"]}+{c}{CF["C"]}', bold=True, fill=FILL_KEY)
    C_('cum', 'Кумулятивный поток наличности', lambda y, c: f'={c}{CF["net"]}' if y == 2027 else f'={L(YC[y] - 1)}{r}+{c}{CF["net"]}',
       bold=True, total='last')
    r += 1
    C_('dscr', 'Коэффициент обслуживания долга, DSCR [(A)/(B+C)]',
       lambda y, c: '="-"' if y == 2027 else f'=IFERROR({c}{CF["A"]}/-({c}{CF["B"]}+{c}{CF["C"]}),0)', unit='x', fmt=MULT, total=None)
    C_('dscr2', 'DSCR с учётом поддерживающих капвложений [(A − инвестиции)/(B+C)]',
       lambda y, c: '="-"' if y == 2027 else f'=IFERROR(({c}{CF["A"]}+{c}{CF["capex"]})/-({c}{CF["B"]}+{c}{CF["C"]}),0)', unit='x', fmt=MULT, total=None)
    put(ws, (r, 2), 'Минимальный DSCR (2–7 годы)', bold=True); put(ws, (r, 4), f'=MIN(E{CF["dscr2"]}:J{CF["dscr2"]})', fmt=MULT, bold=True); CF['dmin'] = r; r += 1
    put(ws, (r, 2), 'Средний DSCR (2–7 годы)', bold=True); put(ws, (r, 4), f'=AVERAGE(E{CF["dscr2"]}:J{CF["dscr2"]})', fmt=MULT, bold=True); CF['davg'] = r; r += 2
    put(ws, (r, 2), 'Проценты льготного периода (1-й год) уплачиваются за счёт собственных средств заемщика (включены в стоимость проекта, Приложение 1). '
        'Налог на прибыль — в размере начисленного (без учёта отсрочки уплаты).', italic=True, border=False)

    # ------------------------------------------------------------------ NPV_IRR
    ws = ws_npv
    title(ws, 'РАСЧЁТ NPV, IRR И СРОКА ОКУПАЕМОСТИ ПРОЕКТА (в экв. долл. США)', last_col=11,
          sub='Денежный поток проекта (до финансирования) = операционный поток (A) − инвестиции; горизонт — срок кредита (7 лет), без остаточной стоимости.')
    year_header(ws, 4, total='Итого')
    r = 6
    N = {}
    def N_(k, label, f, **kw):
        nonlocal r
        N[k] = yrow(ws, r, label, f, **kw)
        r += 1
    N_('t', 'Период (лет)', lambda y, c: YEARS.index(y) + 1, unit='', fmt='0', total=None)
    N_('A', 'Операционный поток наличности (A)', lambda y, c: f'=CashFlow!{c}{CF["A"]}')
    N_('inv', 'Инвестиции', lambda y, c: f'=CashFlow!{c}{CF["capex"]}')
    N_('fcf', 'Поток наличности проекта (FCFF)', lambda y, c: f'={c}{N["A"]}+{c}{N["inv"]}', bold=True, fill=FILL_TOT)
    N_('r', 'Ставка дисконта (WACC)', lambda y, c: f'={src.ref("wacc")}' if y == 2027 else f'=D{r}', unit='%', fmt=PCT, total=None)
    N_('df', 'Коэффициент дисконтирования', lambda y, c: f'=1/(1+{c}{N["r"]})^{c}{N["t"]}', unit='', fmt='0.0000', total=None)
    N_('dcf', 'Дисконтированный поток', lambda y, c: f'={c}{N["fcf"]}*{c}{N["df"]}')
    N_('cum', 'Накопленный поток', lambda y, c: f'={c}{N["fcf"]}' if y == 2027 else f'={L(YC[y] - 1)}{r}+{c}{N["fcf"]}', total='last')
    N_('dcum', 'Накопленный дисконтированный поток (NPV нарастающим итогом)', lambda y, c: f'={c}{N["dcf"]}' if y == 2027 else f'={L(YC[y] - 1)}{r}+{c}{N["dcf"]}', total='last')
    N_('pb', 'вспом.: окупаемость', lambda y, c: '=""' if y == 2027 else
       f'=IF(AND({L(YC[y] - 1)}{N["cum"]}<0,{c}{N["cum"]}>=0),{L(YC[y] - 1)}{N["t"]}-{L(YC[y] - 1)}{N["cum"]}/{c}{N["fcf"]},"")', unit='', fmt='0.00', total=None)
    N_('dpb', 'вспом.: дисконтированная окупаемость', lambda y, c: '=""' if y == 2027 else
       f'=IF(AND({L(YC[y] - 1)}{N["dcum"]}<0,{c}{N["dcum"]}>=0),{L(YC[y] - 1)}{N["t"]}-{L(YC[y] - 1)}{N["dcum"]}/{c}{N["dcf"]},"")', unit='', fmt='0.00', total=None)
    r += 1
    res = [('NPV (чистая приведённая стоимость), USD', f'=SUM(D{N["dcf"]}:J{N["dcf"]})', NUM),
           ('IRR (внутренняя норма доходности)', f'=IFERROR(IRR(D{N["fcf"]}:J{N["fcf"]}),"-")', PCT),
           ('Индекс доходности (PI)', f'=IFERROR(SUM(E{N["dcf"]}:J{N["dcf"]})/-D{N["dcf"]},0)', '0.00'),
           ('Простой срок окупаемости, лет (от начала 1-го года)', f'=IF(COUNT(E{N["pb"]}:J{N["pb"]})>0,SUM(E{N["pb"]}:J{N["pb"]}),"> 7")', '0.0'),
           ('Дисконтированный срок окупаемости, лет', f'=IF(COUNT(E{N["dpb"]}:J{N["dpb"]})>0,SUM(E{N["dpb"]}:J{N["dpb"]}),"> 7")', '0.0')]
    NR = {}
    for lab, f, fm in res:
        put(ws, (r, 2), lab, bold=True, fill=FILL_KEY); put(ws, (r, 3), None, fill=FILL_KEY)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        put(ws, (r, 4), f, fmt=fm, bold=True, fill=FILL_KEY)
        NR[lab[:3]] = r
        r += 1
    NPV_ROW, IRR_ROW, PI_ROW, PB_ROW, DPB_ROW = r - 5, r - 4, r - 3, r - 2, r - 1
    r += 1

    # ------------------------------------------------------------------ Проверки
    ws = ws_chk
    title(ws, 'КОНТРОЛЬНЫЕ ПРОВЕРКИ МОДЕЛИ', last_col=5)
    for i, h in enumerate(['Проверка', 'Значение', 'Допуск', 'Результат']):
        put(ws, (4, 2 + i), h, bold=True, fill=FILL_HEAD, align='center')
    ck = []
    sumabs = lambda a, b: f'SUMPRODUCT(ABS({a}-{b}))'
    rev_model = "'{0}'!F{1}:L{1}".format(SRC, src.rows['rev'])
    ebitda_model = "'{0}'!F{1}:L{1}".format(SRC, src.rows['ebitda'])
    el_model = "'{0}'!F{1}:L{1}".format(SRC, src.rows['el_total'])
    ck += [('Выручка = ТЭО (сумма отклонений за 1–7 годы, USD)', f'={sumabs("ProfLoss!D{0}:J{0}".format(PLR["rev"]), rev_model)}', 1, 'abs'),
           ('EBITDA = ТЭО (сумма отклонений, USD)', f'={sumabs("ProfLoss!D{0}:J{0}".format(PLR["ebitda"]), ebitda_model)}', 1, 'abs'),
           ('Производственная себестоимость по элементам = ТЭО (без амортизации капитализированных процентов)',
            f'=SUMPRODUCT(ABS(CostSold!D{CS["prodcost"]}:J{CS["prodcost"]}-Depreciate!G{DEP["capint"]}:M{DEP["capint"]}-{el_model}))', 1, 'abs'),
           ('Производство = мощность × освоение = ТЭО',
            f'=SUMPRODUCT(ABS(ProdPlan!D{PP["prod0"]}:J{PP["prod0"]}-\'{SRC}\'!F{src.rows["prod0"]}:L{src.rows["prod0"]}))', 0.5, 'abs'),
           ('Запас готовой продукции на конец 7-го года = ТЭО', f'=ABS(ProdPlan!J{PP["stock0"]}-{X("close0", 2033)})', 0.5, 'abs'),
           ('Стоимость проекта = источники финансирования (USD)', f'=ProjCost!E{PCR["tot"]}-FinPlan!E15', 1, 'abs'),
           ('Кредит в финансовом плане = сумма кредита (USD)', f'=FinPlan!E14-{INF["loan"]}', 1, 'abs'),
           ('Кредит полностью погашен к концу 7-го года (остаток, USD)', f'={LNR("close", 2033)}', 1, 'abs'),
           ('Сумма кредитов по трём проектам = 78 000 000 USD', f'=ИНФОРМАЦИЯ!O{r0 + 3}-{LOAN_TOTAL}', 1, 'abs'),
           ('Минимальный чистый поток наличности за 1–7 годы (должен быть > 0)', f'=MIN(CashFlow!D{CF["net"]}:J{CF["net"]})', 0, 'pos'),
           ('Минимальный кумулятивный поток наличности (должен быть > 0)', f'=MIN(CashFlow!D{CF["cum"]}:J{CF["cum"]})', 0, 'pos'),
           ('Минимальная чистая прибыль за 1–7 годы (должна быть > 0)', f'=MIN(ProfLoss!D{PLR["np"]}:J{PLR["np"]})', 0, 'pos'),
           ('Минимальный DSCR за 2–7 годы (≥ 1,20)', f'=CashFlow!D{CF["dmin"]}', 1.2, 'ge'),
           ('Амортизация по объектам = ТЭО (без капитализированных процентов)',
            f'=SUMPRODUCT(ABS(Depreciate!G{DEP["tot"]}:M{DEP["tot"]}-Depreciate!G{DEP["capint"]}:M{DEP["capint"]}-\'{SRC}\'!F{src.rows["da"]}:L{src.rows["da"]}+\'{SRC}\'!F{src.rows["capbc_dep"]}:L{src.rows["capbc_dep"]}))', 1, 'abs'),
           (f'Затраты при проектной мощности (Прил. 10) = себестоимость + расходы периода + прочие ({REF} г., Прил. 11)',
            f'=CostTotal!C{CT["tot"]}-CostSold!{L(YC[REF])}{CS["prodcost"]}-CostSold!{L(YC[REF])}{CS["period"]}-CostSold!{L(YC[REF])}{CS["opx"]}', 1, 'abs'),
           (f'Сырьё, материалы, энергоресурсы, ФОТ и прочие (Прил. 7, 8) = производственная себестоимость без амортизации ({REF} г.)',
            f'=RawMat!F{RM["tot"]}+RawMat!F{RM["other"]}+Labour!F{LB["energy"]}+Labour!F{LB["prod"]}-(CostSold!{L(YC[REF])}{CS["prodcost"]}-CostSold!{L(YC[REF])}{CS["dep"]})', 1, 'abs'),
           (f'Фонд оплаты труда административного персонала с соц. налогом (Прил. 8) = ТЭО ({REF} г.)',
            f'=Labour!F{LB["ga"]}-{X("pers_ga", REF)}', 1, 'abs'),
           ('Поток 1-го года: собственные средства покрывают все затраты сверх кредита (остаток на конец 1-го года ≥ резерва)',
            f'=CashFlow!D{CF["cum"]}-ProjCost!E{PCR["res"]}', 0, 'ge0')]
    r = 5
    for lab, f, tol, kind in ck:
        put(ws, (r, 2), lab, wrap=True)
        put(ws, (r, 3), f, fmt=NUM2 if kind != 'ge' else MULT)
        put(ws, (r, 4), tol, fmt=NUM2 if kind != 'ge' else MULT)
        if kind == 'abs':
            res_f = f'=IF(ABS(C{r})<=D{r},"OK","ОШИБКА")'
        elif kind == 'pos':
            res_f = f'=IF(C{r}>D{r},"OK","ОШИБКА")'
        elif kind == 'ge0':
            res_f = f'=IF(C{r}>=-0.01,"OK","ОШИБКА")'
        else:
            res_f = f'=IF(C{r}>=D{r},"OK","ОШИБКА")'
        put(ws, (r, 5), res_f, bold=True, align='center')
        r += 1
    put(ws, (r + 1, 2), 'ИТОГ ПРОВЕРОК', bold=True, fill=FILL_KEY)
    put(ws, (r + 1, 5), f'=IF(COUNTIF(E5:E{r - 1},"OK")={len(ck)},"ВСЕ OK","ЕСТЬ ОШИБКИ")', bold=True, fill=FILL_KEY, align='center')
    CHK_ROW = r + 1

    # ------------------------------------------------------------------ Резюме
    ws = ws_sum
    put(ws, 'B2', f'БИЗНЕС-ПЛАН ИНВЕСТИЦИОННОГО ПРОЕКТА', bold=True, size=14, border=False)
    put(ws, 'B3', ADV_RU[key][0], bold=True, size=12, border=False)
    put(ws, 'B4', f'({meta["name_ru"]})', italic=True, border=False)
    r = 6
    def sec(t):
        nonlocal r
        put(ws, (r, 2), t, bold=True, fill=FILL_HEAD)
        for c in range(3, 12):
            put(ws, (r, c), None, fill=FILL_HEAD)
        r += 1
    def kv(k, v, fmt=None, bold=False):
        nonlocal r
        put(ws, (r, 2), k, wrap=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=11)
        put(ws, (r, 3), v, fmt=fmt, bold=bold, wrap=True, align='left')
        for c in range(4, 12):
            ws.cell(row=r, column=c).border = BORDER
        if isinstance(v, str) and len(v) > 110:
            ws.row_dimensions[r].height = 30
        r += 1
    sec('1. ОБЩИЕ СВЕДЕНИЯ')
    mk = 'экспорт' if main['market'] == 'export' else 'внутренний рынок'
    kv('Суть проекта', f'{ADV_RU[key][0]}. {ADV_RU[key][2]}; выпуск основной продукции — {main["name_ru"]}, '
                       f'{main["capacity"]:,.1f} {main["unit"]}/год при проектной мощности; реализация — {mk}.'.replace(',', ' ').replace('.0 ', ' '))
    kv('Заемщик', BORROWER)
    kv('Местоположение', meta['site'])
    kv('Продукция', ', '.join(p['name_ru'] + ('' if i == 0 else ' (попутная)') for i, p in enumerate(prods)))
    kv('Производственная мощность', ADV_RU[key][2] + f'; выпуск основной продукции — {main["capacity"]:,.0f} {main["unit"]}/год'.replace(',', ' '))
    kv('Рынок сбыта', 'экспорт (100% выручки в USD)' if main['market'] == 'export' else 'внутренний рынок (поставка концентрата на переработку), цены в USD')
    kv('Сроки', f'начало строительства — {dt.date.fromisoformat(meta["constr_start"]):%d.%m.%Y}; начало производства — {fp:%d.%m.%Y}; '
                f'выход на проектную мощность - {REFL}')
    kv('Стадия проработки', {'Feasibility': 'ТЭО', 'Pre-feasibility': 'Предварительное ТЭО'}.get(meta['stage'], meta['stage']))
    r += 1
    sec('2. ПАРАМЕТРЫ КРЕДИТА')
    kv('Кредитор', BANK)
    kv('Сумма кредита, USD', f'={INF["loan"]}', fmt=NUM, bold=True)
    kv('Процентная ставка', f'={INF["rate"]}', fmt=PCT)
    kv('Срок кредита', '84 месяца (7 лет): льготный (инвестиционный) период - 12 месяцев (1-й год), погашение - 72 месяца (2-7 годы)')
    kv('Порядок погашения', 'основной долг — ежемесячно равными долями; проценты — ежемесячно на остаток задолженности')
    kv('Проценты льготного периода', '=TEXT(' + I_GRACE + ',"#,##0") & " USD — уплачиваются за счёт собственных средств заемщика (включены в стоимость проекта)"')
    kv('Целевое назначение', '30% авансовый платёж по EPC-контракту (E — проектирование, P — поставка оборудования, C — строительство)')
    r += 1
    sec('3. СТОИМОСТЬ ПРОЕКТА И ИСТОЧНИКИ ФИНАНСИРОВАНИЯ')
    put(ws, (r, 2), 'Статья', bold=True, fill=FILL_TOT); put(ws, (r, 3), 'USD', bold=True, fill=FILL_TOT, align='center')
    put(ws, (r, 4), 'Доля', bold=True, fill=FILL_TOT, align='center'); r += 1
    for lab, f in [('Фиксированные активы', f'=ProjCost!E{PCR["fa"]}'), ('Рабочий капитал', f'=ProjCost!E{PCR["wc"]}'),
                   ('Проценты льготного периода и расходы подготовительного периода', f'=ProjCost!E{PCR["fin"]}'),
                   ('Резерв', f'=ProjCost!E{PCR["res"]}'), ('ВСЕГО СТОИМОСТЬ ПРОЕКТА', f'=ProjCost!E{PCR["tot"]}'),
                   ('Собственные средства заемщика', f'={FP_EQ}'), (f'Кредит {BANK_S}', f'={FP_LOAN}')]:
        b = lab.startswith('ВСЕГО')
        put(ws, (r, 2), lab, bold=b, fill=FILL_KEY if b else None); put(ws, (r, 3), f, fmt=NUM, bold=b, fill=FILL_KEY if b else None)
        put(ws, (r, 4), f'=IFERROR(C{r}/ProjCost!$E${PCR["tot"]},0)', fmt=PCT, bold=b, fill=FILL_KEY if b else None)
        r += 1
    r += 1
    sec('4. ОСНОВНЫЕ ФИНАНСОВЫЕ ПОКАЗАТЕЛИ ПО ГОДАМ')
    year_header(ws, r); r += 2
    for lab, f, fm in [('Выручка (без НДС)', lambda y, c: f'=ProfLoss!{c}{PLR["rev"]}', NUM),
                       ('EBITDA', lambda y, c: f'=ProfLoss!{c}{PLR["ebitda"]}', NUM),
                       ('Чистая прибыль', lambda y, c: f'=ProfLoss!{c}{PLR["np"]}', NUM),
                       ('Операционный поток наличности (A)', lambda y, c: f'=CashFlow!{c}{CF["A"]}', NUM),
                       ('Обслуживание долга (проценты + основной долг)', lambda y, c: f'={LNR("ds", y)}', NUM),
                       ('Чистый поток наличности', lambda y, c: f'=CashFlow!{c}{CF["net"]}', NUM),
                       ('Кумулятивный поток наличности', lambda y, c: f'=CashFlow!{c}{CF["cum"]}', NUM),
                       ('DSCR (с учётом поддерживающих капвложений)', lambda y, c: f'=CashFlow!{c}{CF["dscr2"]}', MULT),
                       ('Задолженность по кредиту на конец года', lambda y, c: f'={LNR("close", y)}', NUM)]:
        yrow(ws, r, lab, f, fmt=fm, unit='USD' if fm == NUM else 'x',
             total='sum' if lab.startswith(('Выручка', 'EBITDA', 'Чистая п', 'Операц', 'Обслуж', 'Чистый п')) else None)
        r += 1
    r += 1
    sec('5. ПОКАЗАТЕЛИ ЭФФЕКТИВНОСТИ И УСТОЙЧИВОСТИ')
    for lab, f, fm in [('NPV проекта за 7 лет, USD', f'=NPV_IRR!D{NPV_ROW}', NUM),
                       ('IRR проекта за 7 лет', f'=NPV_IRR!D{IRR_ROW}', PCT),
                       ('Индекс доходности (PI)', f'=NPV_IRR!D{PI_ROW}', '0.00'),
                       ('Простой срок окупаемости, лет', f'=NPV_IRR!D{PB_ROW}', '0.0'),
                       ('Дисконтированный срок окупаемости, лет', f'=NPV_IRR!D{DPB_ROW}', '0.0'),
                       ('Минимальный / средний DSCR (2–7 годы)', f'=TEXT(CashFlow!D{CF["dmin"]},"0.00") & "x / " & TEXT(CashFlow!D{CF["davg"]},"0.00") & "x"', None),
                       ('Точка безубыточности при проектной мощности, % от выручки', '=BreakPoint!C10', PCT)]:
        put(ws, (r, 2), lab); put(ws, (r, 3), f, fmt=fm, bold=True)
        if fm is None:
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        r += 1
    r += 1
    sec('6. ВЫПОЛНЕНИЕ УСЛОВИЙ ФИНАНСИРОВАНИЯ')
    for lab, f in [('Положительный чистый поток наличности в каждом из 1-7 годов', f'=IF(MIN(CashFlow!D{CF["net"]}:J{CF["net"]})>0,"Да","Нет")'),
                   ('Чистая прибыль в каждом из 1-7 годов', f'=IF(MIN(ProfLoss!D{PLR["np"]}:J{PLR["np"]})>0,"Да","Нет")'),
                   ('Погашение кредита в течение 6 лет после льготного периода', f'=IF(ABS({LNR("close", 2033)})<1,"Да","Нет")'),
                   ('Проценты льготного периода покрыты собственными средствами', f'=IF(ProjCost!E{PCR["int"]}>0,"Да","Нет")'),
                   ('Дисконтированный срок окупаемости - в пределах срока кредита (7 лет)', f'=IF(AND(ISNUMBER(NPV_IRR!D{DPB_ROW}),NPV_IRR!D{DPB_ROW}<=7),"Да","Нет")')]:
        put(ws, (r, 2), lab); put(ws, (r, 3), f, bold=True, align='center'); r += 1
    r += 1
    sec('7. ОСНОВНЫЕ ДОПУЩЕНИЯ')
    notes = [
        'Производственные, затратные, налоговые и инвестиционные показатели приняты по ТЭО проекта; цены реализации - на консервативном уровне.',
        f'Кредитор - {BANK}. Сумма кредита определена по 30% авансу по EPC-контрактам (распределение 78 млн USD между тремя проектами - лист «ИНФОРМАЦИЯ»).',
        'Кредит: 7% годовых, 84 месяца, льготный период 12 месяцев, погашение основного долга ежемесячно равными долями; база начисления процентов - 360 дней.',
        'Все затраты сверх суммы кредита финансируются за счёт собственных средств заемщика.',
        'Кредит выдаётся единовременно в начале 1-го года; проценты 1-го года капитализируются в стоимость основных средств.',
        'Резерв на непредвиденные расходы - 1% суммы кредита.',
    ]
    for t in notes:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
        put(ws, (r, 2), '• ' + t, wrap=True, border=False)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    sec('8. ОСНОВНЫЕ РИСКИ И МЕРЫ ПО ИХ СНИЖЕНИЮ')
    main_mat = max(mat_idx, key=lambda j: D['varlines'][j]['cost'][str(REF)]) if mat_idx else None
    risks = [
        ('Ценовой риск', f'="Точка безубыточности при проектной мощности - " & TEXT(BreakPoint!C10,"0%") & " выручки; цены реализации в плане приняты консервативно. '
                         f'Меры: долгосрочные контракты на реализацию, мониторинг рыночных цен."'),
        ('Валютный риск', f'="Доля экспорта в выручке: " & TEXT(SalePlan!{L(YC[REF])}{SP["sexp"]}/SalePlan!{L(YC[REF])}{SP["tot"]},"0%") & '
                          f'"; цены продукции установлены в USD; кредит — в USD (естественное хеджирование валютной позиции)."'),
        ('Риск освоения мощностей', f'="Освоение мощности: 1-й год — " & TEXT(ProdPlan!D{PP["util0"]},"0%") & ", 2-й год — " & '
                                    f'TEXT(ProdPlan!E{PP["util0"]},"0%") & ", 3-й год — " & TEXT(ProdPlan!F{PP["util0"]},"0%") & '
                                    f'"; льготный период 12 мес. и запас DSCR (мин. " & TEXT(CashFlow!D{CF["dmin"]},"0.00") & "x) покрывают период выхода на мощность."'),
        ('Строительный риск', '="Капитальные затраты приняты в фиксированных ценах (EPC-контракт: проектирование, поставка, строительство); '
                              'превышение сметы финансируется за счёт собственных средств заемщика."'),
    ]
    if main_mat is not None and D['varlines'][main_mat]['element'] == 'Raw materials':
        vl = D['varlines'][main_mat]
        risks.append(('Сырьевой риск', f'="Основное сырьё — {vl["name_ru"]}: " & TEXT(RawMat!F{RM["tot"]}/CostTotal!C{CT["tot"]},"0%") & '
                                       f'" затрат при проектной мощности (с учётом прочих материалов). Меры: долгосрочные договоры поставки."'))
    if meta.get('rent_rate'):
        risks.append(('Налоговый риск', f'="Проект является плательщиком специального рентного налога ({meta["rent_rate"]:.0%}) и НДПИ; '
                                        f'налоговая нагрузка учтена в Приложении 9."'))
    for k, v in risks:
        put(ws, (r, 2), k, bold=True, wrap=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=11)
        put(ws, (r, 3), v, wrap=True)
        for c in range(4, 12):
            ws.cell(row=r, column=c).border = BORDER
        ws.row_dimensions[r].height = 30
        r += 1

    # freeze panes / print setup
    for w in (ws_cs, ws_pl, ws_cf, ws_tx, ws_pp, ws_sp, ws_npv):
        w.freeze_panes = 'D6'
    for w in bk.wb.worksheets:
        w.page_setup.orientation = 'landscape'
        w.page_setup.fitToWidth = 1
        w.sheet_properties.pageSetUpPr.fitToPage = True
        w.page_setup.fitToHeight = 0
    bk.wb.save(out_path)
    return REF


if __name__ == '__main__':
    D = json.load(open(sys.argv[1]))
    print('REF year', build(D, sys.argv[2], sys.argv[3]))
