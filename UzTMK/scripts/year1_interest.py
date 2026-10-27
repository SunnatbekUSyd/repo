"""Show the grace-period (year 1) loan interest as an expense of year 1.

CostSold / ProfLoss / CashFlow get the year-1 interest; it stays in the project cost (paid from own funds).
The interest is no longer capitalised into fixed assets (that row is removed from Depreciate, otherwise it
would be expensed twice). Income tax is not negative in a loss year: the loss is carried to the next years.
"""
import re
import sys

import openpyxl
from openpyxl.styles import Font

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
YC = 'DEFGHIJ'


def find(ws, start, col=2):
    rows = [r for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(r, col).value, str) and ws.cell(r, col).value.strip().startswith(start)]
    assert len(rows) == 1, (ws.title, start, rows)
    return rows[0]


# --- CostSold: year-1 interest
cs = wb['CostSold']
ri = find(cs, 'Проценты за кредит')
assert cs.cell(ri, 4).value in (0, None) and cs.cell(ri, 5).value == '=LOANS!$D$106'
cs.cell(ri, 2).value = '   Проценты за кредит (1-й год - за счёт собственных средств заемщика)'
cs.cell(ri, 4).value = '=LOANS!$C$106'

# --- CashFlow: interest of year 1 now in (B); remove it from the other financial costs row
cf = wb['CashFlow']
rf = find(cf, 'Прочие финансовые издержки в инвестиционный период')
assert cf.cell(rf, 4).value == '=-LOANS!$C$106'
cf.cell(rf, 2).value = 'Прочие финансовые издержки в инвестиционный период'
cf.cell(rf, 4).value = 0

# --- ProfLoss note
pl = wb['ProfLoss']
rn = find(pl, '1-й год: проценты за кредит')
pl.cell(rn, 2).value = ('1-й год: проценты за кредит льготного периода отражены в расходах 1-го года и оплачиваются за счёт '
                        'собственных средств заемщика (включены в стоимость проекта, Приложение 1).')

# --- Taxes: no negative tax; a year-1 loss reduces the taxable profit of the following years
tx = wb['Taxes']
rb, rr, rt = find(tx, 'Налогооблагаемая прибыль'), find(tx, 'Ставка налога на прибыль'), [r for r in range(1, tx.max_row + 1) if tx.cell(r, 2).value == 'Налог на прибыль'][0]
for c in YC:
    assert tx[f'{c}{rt}'].value == f'={c}{rb}*{c}{rr}', tx[f'{c}{rt}'].value
    tx[f'{c}{rt}'].value = f'=MAX(0,MIN({c}{rb},SUM($D{rb}:{c}{rb})))*{c}{rr}'

# --- Резюме assumption text
rs = wb['Резюме']
for row in rs.iter_rows():
    for c in row:
        if isinstance(c.value, str) and 'проценты 1-го года капитализируются' in c.value:
            c.value = c.value.replace('проценты 1-го года капитализируются в стоимость основных средств',
                                      'проценты 1-го года отражаются в расходах 1-го года и оплачиваются за счёт собственных средств заемщика')

# --- Depreciate: remove the capitalised-interest row
dp = wb['Depreciate']
R = find(dp, 'Капитализированные проценты')
rtot = find(dp, 'ВСЕГО')
assert rtot == R + 1
dp['B3'].value = 'Линейный метод.'
for c in ['D'] + list('GHIJKLM'):
    f = dp[f'{c}{rtot}'].value
    assert f.endswith(f':{c}{R})'), f
    dp[f'{c}{rtot}'].value = f[:-len(f'{R})')] + f'{R - 1})'
rnbv = find(dp, 'Остаточная стоимость')
f = dp[f'G{rnbv}'].value
assert f'+$D${R}' in f, f
dp[f'G{rnbv}'].value = f.replace(f'+$D${R}', '')

heights = {r: dp.row_dimensions[r].height for r in range(R + 1, dp.max_row + 1)}
dp.delete_rows(R, 1)
for r, h in heights.items():
    dp.row_dimensions[r - 1].height = h

TOK = re.compile(r"(?<![A-Za-z0-9_.А-Яа-я])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(\$?[A-Z]{1,3}\$?)(\d+)(?![\d(])")


def shift(formula, own):
    def f(m):
        pre, col, row = m.group(1), m.group(2), int(m.group(3))
        if not ((pre is None and own) or pre in ('Depreciate!', "'Depreciate'!")) or row < R:
            return m.group(0)
        assert row > R, ('reference to the removed row', formula)
        return f'{pre or ""}{col}{row - 1}'
    parts = re.split(r'("[^"]*")', formula)
    return ''.join(p if p.startswith('"') else TOK.sub(f, p) for p in parts)


n = 0
for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                v = shift(c.value, w.title == 'Depreciate')
                if v != c.value:
                    c.value, n = v, n + 1
wb.save(out)
print(src.split('/')[-1], 'interest row', ri, 'tax row', rt, 'removed Depreciate row', R, 'refs shifted', n)
